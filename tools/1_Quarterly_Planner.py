import streamlit as st
import pandas as pd
import plotly.express as px
import re
import os
import io
import html
import base64
from dotenv import load_dotenv
from utils.jira_helpers import build_quarterly_epic_progress_table

# Shared file paths for live collaboration
SHARED_DIR = "data"
SHARED_BACKLOG_PATH = os.path.join(SHARED_DIR, "shared_backlog.csv")
SHARED_SPRINT_CALENDAR_PATH = os.path.join(SHARED_DIR, "shared_sprint_calendar.csv")
os.makedirs(SHARED_DIR, exist_ok=True)

def sync_from_disk_if_modified():
    """Reads modifications from the shared disk files if another user changed them."""
    # Sync main backlog
    if os.path.exists(SHARED_BACKLOG_PATH):
        mtime = os.path.getmtime(SHARED_BACKLOG_PATH)
        if 'main_df' not in st.session_state or mtime > st.session_state.get('last_backlog_mtime', 0):
            try:
                df = pd.read_csv(SHARED_BACKLOG_PATH)
                df['Start Date'] = pd.to_datetime(df['Start Date'], errors='coerce')
                df['Due Date'] = pd.to_datetime(df['Due Date'], errors='coerce')
                st.session_state.main_df = df
                st.session_state.last_backlog_mtime = mtime
                st.session_state.editor_key_counter = st.session_state.get('editor_key_counter', 0) + 1
                # Reset selected epic key if no longer present in dataset
                if st.session_state.get('selected_epic_key') not in df['Key'].values:
                    st.session_state.selected_epic_key = None
            except Exception:
                pass

    # Sync sprint calendar
    if os.path.exists(SHARED_SPRINT_CALENDAR_PATH):
        mtime_sprints = os.path.getmtime(SHARED_SPRINT_CALENDAR_PATH)
        if 'sprint_calendar' not in st.session_state or mtime_sprints > st.session_state.get('last_sprint_cal_mtime', 0):
            try:
                df_sprints = pd.read_csv(SHARED_SPRINT_CALENDAR_PATH)
                st.session_state.sprint_calendar = df_sprints
                st.session_state.last_sprint_cal_mtime = mtime_sprints
            except Exception:
                pass

# Run initial disk sync check
sync_from_disk_if_modified()

# Load environment variables
load_dotenv(override=True)
default_primary_color = os.getenv("PRIMARY_COLOR", "#3B82F6")
if 'primary_color' not in st.session_state or st.session_state.get('prev_env_color') != default_primary_color:
    st.session_state.primary_color = default_primary_color
    st.session_state.prev_env_color = default_primary_color


if 'selected_epic_key' not in st.session_state:
    st.session_state.selected_epic_key = None
if 'last_gantt_selection' not in st.session_state:
    st.session_state.last_gantt_selection = None
if 'chart_key_counter' not in st.session_state:
    st.session_state.chart_key_counter = 0
if 'editor_key_counter' not in st.session_state:
    st.session_state.editor_key_counter = 0
if 'just_reset_selection' not in st.session_state:
    st.session_state.just_reset_selection = False
if 'quarterly_progress_df' not in st.session_state:
    st.session_state.quarterly_progress_df = None
if 'quarterly_progress_config' not in st.session_state:
    st.session_state.quarterly_progress_config = {}

# Custom CSS for Premium Dark Mode Theme
st.markdown("""
    <style>
        /* General dark theme override */
        .main, .block-container, [data-testid="stAppViewContainer"], [data-testid="stHeader"] {
            background-color: #27272E !important;
            color: #F8FAFC !important;
        }
        
        /* High visibility and premium SaaS styling for file uploader */
        [data-testid="stFileUploadDropzone"] {
            background-color: #18181D !important;
            border: 1px dashed #60A5FA !important;
            border-radius: 12px !important;
        }
        [data-testid="stFileUploadDropzone"] * {
            color: #F8FAFC !important;
        }
        [data-testid="stFileUploaderFileName"], 
        .stFileUploaderFileName, 
        .uploadedFileName {
            color: #FFFFFF !important;
            font-weight: 600 !important;
        }
        [data-testid="stFileUploader"] section {
            background-color: #18181D !important;
            border: 1px solid #3E3E4A !important;
            border-radius: 12px !important;
        }
        [data-testid="stFileUploader"] section * {
            color: #FFFFFF !important;
        }
        
        /* Typography high-contrast styles */
        h1, h2, h3, h4, h5, h6 {
            color: #FFFFFF !important;
            font-weight: 700 !important;
        }
        .stMarkdown p, .stMarkdown li, .stMarkdown span {
            color: #F8FAFC !important;
        }
        label, .stWidgetLabel {
            color: #FFFFFF !important;
            font-weight: 600 !important;
        }
        
        /* Sidebar Toggle Button */
        button[data-testid="collapsedControl"], button[kind="header"] {
            color: #60A5FA !important;
        }
        button[data-testid="collapsedControl"] svg, button[kind="header"] svg {
            fill: #60A5FA !important;
        }
        
        /* File Uploader small limit text */
        [data-testid="stFileUploadDropzone"] small {
            color: #94A3B8 !important;
            font-weight: 500 !important;
        }
        
        /* Sidebar dark theme label improvements */
        section[data-testid="stSidebar"] {
            background-color: #18181D !important;
            border-right: 1px solid #3E3E4A !important;
        }
        section[data-testid="stSidebar"] h1,
        section[data-testid="stSidebar"] h2,
        section[data-testid="stSidebar"] h3,
        section[data-testid="stSidebar"] h4,
        section[data-testid="stSidebar"] label,
        section[data-testid="stSidebar"] .stMarkdown {
            color: #FFFFFF !important;
        }
        
        /* Rounded borders and shadows for dataframes and tables */
        div[data-testid="stDataFrame"] { 
            border-radius: 12px; 
            overflow: hidden; 
            border: 1px solid #3E3E4A !important;
            box-shadow: 0 10px 15px -3px rgba(0, 0, 0, 0.5);
            background-color: #18181D;
        }
        div[data-testid="stExpander"] { 
            border-radius: 12px !important; 
            border: 1px solid #3E3E4A !important; 
            background-color: #18181D !important;
            box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.4);
            margin-bottom: 1rem;
            overflow: hidden;
        }
        div[data-testid="stExpander"] > details {
            background-color: #18181D !important;
        }
        div[data-testid="stExpander"] summary {
            background-color: #18181D !important;
            color: #FFFFFF !important;
            font-weight: 600;
        }
        div[data-testid="stExpander"] summary:hover {
            background-color: #2D2D38 !important;
        }
        div[data-testid="stExpander"] .stMarkdown {
            color: #F8FAFC !important;
        }
        
        /* General bordered container style */
        div[data-testid="stVerticalBlockBorderWrapper"],
        div.stVerticalBlockBorder {
            border: 1px solid #3E3E4A !important;
            border-radius: 12px !important;
            background-color: #18181D !important;
            padding: 1.5rem !important;
            box-shadow: 0 10px 15px -3px rgba(0, 0, 0, 0.5) !important;
        }
        
        /* Gantt chart container box specifically targeted with key or contents */
        div[data-testid="stVerticalBlockBorderWrapper"]:has(.stPlotlyChart),
        div[data-testid="stVerticalBlockBorderWrapper"]:has([data-testid="stPlotlyChart"]),
        div[data-testid="stVerticalBlockBorderWrapper"]:has(.js-plotly-plot),
        div[data-testid="stVerticalBlockBorderWrapper"]:has(iframe),
        div[data-testid="stVerticalBlockBorderWrapper"]:has(.st-key-gantt_container),
        div[data-testid="stVerticalBlockBorderWrapper"]:has([key="gantt_container"]),
        div[data-testid="stVerticalBlockBorderWrapper"]:has([data-testid="st-key-gantt_container"]),
        div[data-testid="stVerticalBlockBorderWrapper"]:has([class*="gantt_container"]),
        div.stVerticalBlockBorder:has(.stPlotlyChart),
        div.stVerticalBlockBorder:has([data-testid="stPlotlyChart"]),
        .st-key-gantt_container {
            background-color: #18181D !important; /* Cohesive Charcoal background */
            border: 2px solid #3E3E4A !important; /* High contrast clean border */
            box-shadow: 0 10px 25px -5px rgba(0, 0, 0, 0.6) !important;
            border-radius: 12px !important;
            padding: 1.5rem !important;
        }
        
        /* Premium button styles with interactive glow and forced high-contrast text */
        .stButton>button, .stDownloadButton>button, button[data-testid="stBaseButton-secondary"] { 
            border-radius: 24px; 
            border: none; 
            background: linear-gradient(135deg, #3B82F6 0%, #1D4ED8 100%) !important; 
            color: #FFFFFF !important; 
            padding: 0.5rem 1.5rem; 
            font-weight: 600;
            box-shadow: none !important;
            transition: all 0.2s ease-in-out;
        }
        .stButton>button:hover, .stDownloadButton>button:hover, button[data-testid="stBaseButton-secondary"]:hover { 
            background: linear-gradient(135deg, #2563EB 0%, #1D4ED8 100%) !important;
            color: #FFFFFF !important;
            box-shadow: none !important;
            transform: translateY(-1px);
        }
        
        /* Quality validation alert card with clean high contrast text */
        .validation-card {
            background-color: #1E1B4B;
            border-left: 4px solid #F59E0B;
            padding: 1rem;
            border-radius: 8px;
            margin-bottom: 1rem;
            color: #F8FAFC;
        }
    </style>
""", unsafe_allow_html=True)

# Helper to parse Quarter (e.g. "Q1 - 26") to a starting Date
def parse_quarter_to_date(q_str):
    if pd.isna(q_str) or not isinstance(q_str, str): return pd.Timestamp.today()
    match = re.search(r'Q([1-4]).*?(\d{2,4})', q_str, re.IGNORECASE)
    if match:
        q = int(match.group(1))
        year = int(match.group(2))
        if year < 100: year += 2000
        month = (q - 1) * 3 + 1
        return pd.Timestamp(year=year, month=month, day=1)
    return pd.Timestamp.today()


def get_status_group(status_str, is_milestone=False):
    """
    Unifies status mapping across the Gantt chart and selected detail panel views.
    """
    if is_milestone:
        return "Milestone"
    s = str(status_str).strip().lower()
    if s in ["to do", "open", "backlog", "pending", "por hacer", "abierto"]:
        return "To Do"
    elif s in ["in progress", "en progreso", "desarrollo", "testing", "in_progress"]:
        return "In Progress"
    elif s in ["done", "completado", "listo", "closed", "cerrado", "finalizado", "terminado", "acceptance test"]:
        return "Done"
    elif s in ["blocked", "bloqueado", "impedimento", "pausado"]:
        return "Blocked"
    return "To Do"


def truncate_text(text, max_len=45):
    """
    Cleans and truncates text values for optimized rendering inside Plotly visualizations.
    """
    if pd.isna(text):
        return ""
    text_str = str(text)
    return text_str[:max_len-3] + "..." if len(text_str) > max_len else text_str


def clean_val(val, fallback="N/A"):
    """
    Sanitizes values to avoid rendering empty or 'NaN' elements in active UI layouts.
    """
    if pd.isna(val) or val == "" or str(val).strip().lower() == "nan":
        return fallback
    return str(val)


def make_select_option(r):
    """
    Formats a dictionary row representing a Jira epic into a clean dropdown selector option.
    """
    k = clean_val(r.get('Key'), 'N/A')
    n = clean_val(r.get('Epic Name'), 'Unnamed Epic')
    return f"[{k}] {n}"


def val_changed(v1, v2):
    """
    Safe comparison check preventing TypeError: 'boolean value of NA is ambiguous' errors in pandas.
    """
    n1, n2 = pd.isna(v1), pd.isna(v2)
    if n1 and n2:
        return False
    if n1 != n2:
        return True
    try:
        return bool(v1 != v2)
    except Exception:
        return True


def get_quarterly_team_label(labels):
    """Return the first configured team label assigned to an Epic."""
    configured_labels = [label.strip() for label in os.getenv("TEAM_LABELS", "").split(",") if label.strip()]
    item_labels = {label.strip().lower() for label in re.split(r'[\s,]+', str(labels)) if label.strip()}
    return next((label for label in configured_labels if label.lower() in item_labels), "-")


def prepare_quarterly_progress_for_presentation(df):
    """Add presentation fields and prioritize open, lower-progress Epics for review."""
    prepared = df.copy()
    presentation_defaults = {
        "Presentation update": "",
        "Include in delivery roadmap": False,
        "Release version": "",
        "Delivery month": "",
        "Delivery timing": "Mid",
    }
    for column, default_value in presentation_defaults.items():
        if column not in prepared.columns:
            prepared[column] = default_value
    prepared["Team"] = prepared.get("Labels", pd.Series("", index=prepared.index)).apply(get_quarterly_team_label)

    configured_labels = [label.strip() for label in os.getenv("TEAM_LABELS", "").split(",") if label.strip()]
    priorities = {label.lower(): position for position, label in enumerate(configured_labels)}
    prepared["_team_order"] = prepared["Team"].str.lower().map(priorities).fillna(len(priorities))
    status_values = prepared.get("Status", pd.Series("", index=prepared.index)).fillna("").astype(str).str.strip().str.lower()
    # Quarterly progress maps a closed Epic to Done.  Keep accepted variants for
    # existing imported data, then place these Epics after all open work.
    prepared["_is_closed_epic"] = status_values.isin({"done", "closed", "resolved"}).astype(int)
    prepared["_completion_order"] = pd.to_numeric(
        prepared.get("Completion", pd.Series("0", index=prepared.index))
        .fillna("0")
        .astype(str)
        .str.replace("%", "", regex=False),
        errors="coerce",
    ).fillna(0)
    prepared = prepared.sort_values(
        ["_is_closed_epic", "_team_order", "_completion_order", "Key"],
        kind="stable",
    ).drop(columns=["_is_closed_epic", "_team_order", "_completion_order"])
    return prepared.reset_index(drop=True)


def delivery_month_options():
    """Provide a simple rolling set of month choices for approximate delivery milestones."""
    start = pd.Timestamp.today().replace(day=1)
    return ["-"] + [(start + pd.DateOffset(months=offset)).strftime("%b %Y") for offset in range(18)]


def build_delivery_roadmap_timeline(df):
    """Turn selected delivery commitments into an easy-to-read milestone timeline."""
    roadmap = df[df["Include in delivery roadmap"].fillna(False).astype(bool)].copy()
    roadmap = roadmap[
        roadmap["Delivery month"].fillna("").astype(str).str.strip().ne("")
        & roadmap["Delivery month"].fillna("").astype(str).str.strip().ne("-")
    ]
    if roadmap.empty:
        return roadmap, None

    timing_days = {"Beginning": 3, "Mid": 15, "End": 25}
    def milestone_date(row):
        try:
            month_start = pd.to_datetime(f"01 {row['Delivery month']}", format="%d %b %Y")
            return month_start + pd.Timedelta(days=timing_days.get(str(row.get("Delivery timing", "Mid")), 15) - 1)
        except (TypeError, ValueError):
            return pd.NaT

    roadmap["Milestone date"] = roadmap.apply(milestone_date, axis=1)
    roadmap = roadmap.dropna(subset=["Milestone date"]).sort_values("Milestone date", kind="stable")
    if roadmap.empty:
        return roadmap, None

    roadmap["Epic"] = roadmap["Summary"].fillna("Unnamed Epic").astype(str)
    roadmap["Release"] = roadmap["Release version"].fillna("").replace("", "Release TBD")
    roadmap["Update"] = roadmap["Presentation update"].fillna("").replace("", "No additional context provided")
    fig = px.scatter(
        roadmap,
        x="Milestone date",
        y="Epic",
        symbol="Delivery timing",
        text="Release",
        custom_data=["Update", "Completion", "Status"],
    )
    fig.add_vline(x=pd.Timestamp.today(), line_width=2, line_dash="dash", line_color="#EF4444")
    fig.add_annotation(
        x=pd.Timestamp.today(), y=1.03, yref="paper", text="<b>Today</b>", showarrow=False,
        font=dict(color="#EF4444", size=11), xanchor="center"
    )
    fig.update_traces(
        marker=dict(size=17, line=dict(width=2, color="#FFFFFF")),
        textposition="top center",
        textfont=dict(size=11, color="#F8FAFC"),
        hovertemplate=(
            "<b>%{y}</b><br>"
            "<b>Target:</b> %{x|%d %b %Y}<br>"
            "<b>Release:</b> %{text}<br>"
            "<b>Completion:</b> %{customdata[1]}<br>"
            "<b>Status:</b> %{customdata[2]}<br>"
            "<b>Context:</b> %{customdata[0]}<extra></extra>"
        ),
    )
    fig.update_layout(
        title="Planned delivery milestones",
        xaxis_title="", yaxis_title="", height=max(380, 115 + len(roadmap) * 68),
        plot_bgcolor="#18181D", paper_bgcolor="#18181D", font=dict(color="#F8FAFC"),
        legend_title_text="Approx. timing", margin=dict(l=20, r=20, t=65, b=30),
    )
    fig.update_xaxes(showgrid=True, gridcolor="#475569", tickformat="%b\n%Y")
    fig.update_yaxes(showgrid=False, autorange="reversed")
    return roadmap, fig


def build_delivery_roadmap_slide_pdf(roadmap_df, title, primary_color_hex):
    """Create a compact client-facing PDF slide with delivery milestones on a time axis."""
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib import colors
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    width, height = landscape(letter)
    pdf = canvas.Canvas(buffer, pagesize=(width, height))
    primary = colors.HexColor(primary_color_hex)
    background = colors.HexColor("#F8FAFC")
    dark = colors.HexColor("#1E293B")
    muted = colors.HexColor("#64748B")

    pdf.setFillColor(background)
    pdf.rect(0, 0, width, height, stroke=0, fill=1)
    pdf.setFillColor(primary)
    pdf.rect(0, height - 18, width, 18, stroke=0, fill=1)
    pdf.rect(0, 0, 8, height, stroke=0, fill=1)
    pdf.rect(width - 8, 0, 8, height, stroke=0, fill=1)

    pdf.setFillColor(primary)
    pdf.setFont("Helvetica-Bold", 24)
    pdf.drawString(54, height - 58, str(title)[:85])
    pdf.setFillColor(muted)
    pdf.setFont("Helvetica", 10)
    pdf.drawString(54, height - 76, "Committed Epic delivery roadmap - planned milestones for unfinished work")

    if roadmap_df.empty:
        pdf.setFillColor(dark)
        pdf.setFont("Helvetica", 14)
        pdf.drawString(54, height - 130, "No delivery milestones selected.")
        pdf.save()
        buffer.seek(0)
        return buffer.getvalue()

    earliest = roadmap_df["Milestone date"].min().replace(day=1)
    latest = (roadmap_df["Milestone date"].max().replace(day=1) + pd.DateOffset(months=1))
    if latest <= earliest:
        latest = earliest + pd.DateOffset(months=1)
    total_days = max((latest - earliest).days, 1)
    axis_left, axis_right, axis_y = 280, width - 58, height - 130

    def x_for_date(value):
        return axis_left + ((value - earliest).days / total_days) * (axis_right - axis_left)

    pdf.setStrokeColor(colors.HexColor("#94A3B8"))
    pdf.setLineWidth(2)
    pdf.line(axis_left, axis_y, axis_right, axis_y)
    month = earliest
    while month <= latest:
        x_pos = x_for_date(month)
        pdf.setStrokeColor(colors.HexColor("#CBD5E1"))
        pdf.setLineWidth(0.8)
        pdf.line(x_pos, 66, x_pos, axis_y + 10)
        pdf.setFillColor(muted)
        pdf.setFont("Helvetica-Bold", 8)
        pdf.drawCentredString(x_pos, axis_y + 16, month.strftime("%b %Y"))
        month += pd.DateOffset(months=1)

    today_x = x_for_date(pd.Timestamp.today())
    if axis_left <= today_x <= axis_right:
        pdf.setStrokeColor(colors.HexColor("#EF4444"))
        pdf.setDash(4, 3)
        pdf.line(today_x, 58, today_x, axis_y + 28)
        pdf.setDash()
        pdf.setFillColor(colors.HexColor("#EF4444"))
        pdf.setFont("Helvetica-Bold", 8)
        pdf.drawCentredString(today_x, axis_y + 30, "TODAY")

    available_height = axis_y - 70
    row_height = min(46, max(25, available_height / max(len(roadmap_df), 1)))
    max_rows = int(available_height // row_height)
    display_rows = roadmap_df.head(max_rows)
    for index, (_, row) in enumerate(display_rows.iterrows()):
        y_pos = axis_y - 33 - index * row_height
        marker_x = x_for_date(row["Milestone date"])
        pdf.setStrokeColor(colors.HexColor("#CBD5E1"))
        pdf.setLineWidth(0.5)
        pdf.line(54, y_pos - 13, axis_right, y_pos - 13)
        from reportlab.lib.utils import simpleSplit
        pdf.setFillColor(dark)
        pdf.setFont("Helvetica-Bold", 7.5)
        epic_label = str(row.get("Summary", "Unnamed Epic"))
        epic_lines = simpleSplit(epic_label, "Helvetica-Bold", 7.5, 215)[:3]
        for line_index, line in enumerate(epic_lines):
            pdf.drawString(54, y_pos + 6 - line_index * 8, line)
        release = row.get("Release", "Release TBD")
        pdf.setStrokeColor(primary)
        pdf.setLineWidth(1.2)
        pdf.line(marker_x, y_pos + 6, marker_x, axis_y)
        pdf.setFillColor(primary)
        pdf.circle(marker_x, y_pos + 6, 5, stroke=0, fill=1)
        pdf.setFillColor(muted)
        pdf.setFont("Helvetica-Bold", 7.5)
        pdf.drawString(marker_x + 8, y_pos + 3, str(release))

    if len(roadmap_df) > len(display_rows):
        pdf.setFillColor(muted)
        pdf.setFont("Helvetica-Oblique", 8)
        pdf.drawRightString(axis_right, 46, f"+ {len(roadmap_df) - len(display_rows)} additional milestone(s) shown in the interactive timeline")
    pdf.setFillColor(muted)
    pdf.setFont("Helvetica", 8)
    pdf.drawString(54, 32, f"Generated {pd.Timestamp.today().strftime('%d %b %Y')} | Delivery dates are approximate (Beginning, Mid, End of month).")
    pdf.save()
    buffer.seek(0)
    return buffer.getvalue()


def build_quarterly_progress_slide_pdf(df, title, primary_color_hex):
    """Create a presentation-style landscape slide for the quarterly Epic progress."""
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle
    from utils.pdf_helpers import get_progress_bar_drawing

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(letter),
        leftMargin=54, rightMargin=54, topMargin=48, bottomMargin=42
    )
    styles = getSampleStyleSheet()
    primary = colors.HexColor(primary_color_hex)
    title_style = ParagraphStyle(
        "QuarterlyProgressTitle", parent=styles["Normal"], fontName="Helvetica-Bold",
        fontSize=25, leading=30, textColor=primary, spaceAfter=5
    )
    subtitle_style = ParagraphStyle(
        "QuarterlyProgressSubtitle", parent=styles["Normal"], fontName="Helvetica",
        fontSize=10, leading=13, textColor=colors.HexColor("#475569"), spaceAfter=16
    )
    header_style = ParagraphStyle(
        "QuarterlyProgressHeader", parent=styles["Normal"], fontName="Helvetica-Bold",
        fontSize=8.5, leading=10, textColor=colors.white
    )
    body_style = ParagraphStyle(
        "QuarterlyProgressBody", parent=styles["Normal"], fontName="Helvetica",
        fontSize=8, leading=10, textColor=colors.HexColor("#1E293B")
    )
    key_style = ParagraphStyle("QuarterlyProgressKey", parent=body_style, fontName="Helvetica-Bold")

    visible_columns = [
        ("Key", 55), ("Epic", 150), ("Progress", 60), ("Closed / Resolved", 56),
        ("To Do", 45), ("In progress", 58), ("Status", 62), ("Team", 75), ("Update", 123)
    ]
    column_sources = {"Epic": "Summary", "Progress": "Completion", "Update": "Presentation update"}
    table_data = [[Paragraph(name, header_style) for name, _ in visible_columns]]
    for _, row in df.iterrows():
        cells = []
        for name, _ in visible_columns:
            source = column_sources.get(name, name)
            value = row.get(source, "-")
            value = "-" if pd.isna(value) or str(value).strip() == "" else str(value)
            if name == "Progress":
                cells.append(get_progress_bar_drawing(value, primary, width=52, height=18) or Paragraph(value, body_style))
            else:
                cells.append(Paragraph(html.escape(value), key_style if name == "Key" else body_style))
        table_data.append(cells)

    table = Table(table_data, colWidths=[width for _, width in visible_columns], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), primary),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
    ]))
    project_name = os.getenv("PROJECT_NAME", "RECALL2")
    generated = pd.Timestamp.today().strftime("%d %b %Y")
    story = [
        Paragraph(html.escape(title), title_style),
        Paragraph(f"<b>Project:</b> {html.escape(project_name)} &nbsp; | &nbsp; <b>Generated:</b> {generated} &nbsp; | &nbsp; <b>Committed Epics:</b> {len(df)}", subtitle_style),
        table,
    ]
    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def build_quarterly_progress_pptx(progress_df, roadmap_df, progress_title, progress_subtitle, roadmap_title, roadmap_subtitle, primary_color_hex):
    """Build editable PowerPoint slides for quarterly progress and delivery milestones."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.dml.color import RGBColor

    def rgb(hex_value):
        value = hex_value.lstrip("#")
        return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))

    primary = rgb(primary_color_hex)
    navy = RGBColor(11, 39, 86)
    dark = RGBColor(30, 41, 59)
    muted = RGBColor(100, 116, 139)
    light = RGBColor(241, 245, 249)
    grid = RGBColor(203, 213, 225)
    green = RGBColor(34, 197, 94)
    template_background = RGBColor(0, 45, 55)
    template_teal = RGBColor(0, 151, 143)
    white = RGBColor(255, 255, 255)
    template_muted = RGBColor(113, 190, 187)

    template_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "assets", "templates", "quarterly_planning_template.pptx")
    using_template = os.path.exists(template_path)
    prs = Presentation(template_path) if using_template else Presentation()
    if using_template:
        for slide_id in list(prs.slides._sldIdLst):
            prs.part.drop_rel(slide_id.rId)
            prs.slides._sldIdLst.remove(slide_id)
    else:
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    content_layout = prs.slide_layouts[2] if using_template and len(prs.slide_layouts) > 2 else blank_layout

    def add_text(slide, text, left, top, width, height, size=12, color=dark, bold=False, align=PP_ALIGN.LEFT):
        shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        text_frame = shape.text_frame
        text_frame.clear()
        text_frame.word_wrap = True
        text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        paragraph = text_frame.paragraphs[0]
        paragraph.alignment = align
        run = paragraph.add_run()
        run.text = str(text)
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        return shape

    def add_chrome(slide, slide_title, subtitle=""):
        if using_template:
            # The corporate layout contains editable title placeholders.  They are
            # removed here because python-pptx does not preserve their type styling
            # reliably, causing duplicate, overlapping title text in the export.
            for shape in list(slide.placeholders):
                element = shape._element
                element.getparent().remove(element)
            add_text(slide, slide_title, 0.42, 0.62, 11.6, 0.46, size=24, color=white, bold=False)
            if subtitle:
                add_text(slide, subtitle, 0.43, 1.17, 11.4, 0.27, size=13, color=white)
            return
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(248, 250, 252)
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.18))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = primary
        top_bar.line.fill.background()
        for x_pos in [0, 13.25]:
            band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x_pos), 0, Inches(0.08), prs.slide_height)
            band.fill.solid()
            band.fill.fore_color.rgb = primary
            band.line.fill.background()
        add_text(slide, slide_title, 0.7, 0.35, 11.9, 0.46, size=25, color=navy, bold=True)
        if subtitle:
            add_text(slide, subtitle, 0.7, 0.82, 11.7, 0.24, size=10, color=muted)

    progress_columns = [
        ("Key", "Key", 0.85), ("Epic", "Summary", 2.70), ("Progress", "Completion", 0.85),
        ("Closed / Resolved", "Closed / Resolved", 0.90), ("To Do", "To Do", 0.65),
        ("In progress", "In progress", 0.90), ("Status", "Status", 0.90), ("Team", "Team", 1.10),
        ("Update", "Presentation update", 3.13),
    ]
    rows_per_slide = 12
    progress_df = prepare_quarterly_progress_for_presentation(progress_df)
    for start in range(0, len(progress_df), rows_per_slide):
        chunk = progress_df.iloc[start:start + rows_per_slide].reset_index(drop=True)
        suffix = "" if start == 0 else f" (cont. {start // rows_per_slide + 1})"
        slide = prs.slides.add_slide(content_layout)
        add_chrome(slide, f"{progress_title}{suffix}", progress_subtitle)
        table_left, table_top, table_width = 0.42, 1.90, 12.45
        table_height = 4.92 if using_template else 5.85
        row_height = table_height / (len(chunk) + 1)
        table_shape = slide.shapes.add_table(len(chunk) + 1, len(progress_columns), Inches(table_left), Inches(table_top), Inches(table_width), Inches(table_height))
        table = table_shape.table
        for row in table.rows:
            row.height = Inches(row_height)
        for index, (_, _, column_width) in enumerate(progress_columns):
            table.columns[index].width = Inches(column_width)
            cell = table.cell(0, index)
            cell.fill.solid()
            cell.fill.fore_color.rgb = template_teal if using_template else navy
            cell.text = progress_columns[index][0]
            paragraph = cell.text_frame.paragraphs[0]
            paragraph.runs[0].font.name = "Arial"
            paragraph.runs[0].font.size = Pt(8)
            paragraph.runs[0].font.bold = True
            paragraph.runs[0].font.color.rgb = white
        for row_index, (_, row) in enumerate(chunk.iterrows(), start=1):
            for col_index, (_, source, _) in enumerate(progress_columns):
                cell = table.cell(row_index, col_index)
                cell.fill.solid()
                cell.fill.fore_color.rgb = template_background if using_template else (RGBColor(255, 255, 255) if row_index % 2 else light)
                cell.margin_left = Inches(0.05)
                cell.margin_right = Inches(0.05)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                value = row.get(source, "-")
                value = "-" if pd.isna(value) or str(value).strip() == "" else str(value)
                cell.text = "" if source == "Completion" else value
                if source != "Completion":
                    paragraph = cell.text_frame.paragraphs[0]
                    paragraph.runs[0].font.name = "Arial"
                    paragraph.runs[0].font.size = Pt(7.2)
                    paragraph.runs[0].font.color.rgb = white if using_template else dark
                    if source == "Key":
                        paragraph.runs[0].font.bold = True

                if source == "Completion":
                    try:
                        percentage = max(0, min(100, int(float(value.replace("%", "")))))
                    except Exception:
                        percentage = 0
                    # Keep the percentage and its visual bar *inside* the table
                    # cell.  Separate shapes drift when PowerPoint recalculates
                    # table-row heights, while text runs remain attached to the row.
                    text_frame = cell.text_frame
                    text_frame.clear()
                    text_frame.word_wrap = False
                    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                    label = text_frame.paragraphs[0]
                    label.alignment = PP_ALIGN.CENTER
                    label_run = label.add_run()
                    label_run.text = f"{percentage}%"
                    label_run.font.name = "Arial"
                    label_run.font.size = Pt(6.5)
                    label_run.font.bold = True
                    label_run.font.color.rgb = white if using_template else dark
                    bar = text_frame.add_paragraph()
                    bar.alignment = PP_ALIGN.CENTER
                    # Five compact segments fit in the narrow Progress column in
                    # PowerPoint without wrapping onto a second line.
                    filled_blocks = int(round(percentage / 20))
                    filled_run = bar.add_run()
                    filled_run.text = "▬" * filled_blocks
                    filled_run.font.name = "Arial"
                    filled_run.font.size = Pt(7)
                    filled_run.font.color.rgb = RGBColor(170, 255, 0) if using_template and percentage == 100 else (template_teal if using_template else (green if percentage == 100 else navy))
                    remaining_run = bar.add_run()
                    remaining_run.text = "▬" * (5 - filled_blocks)
                    remaining_run.font.name = "Arial"
                    remaining_run.font.size = Pt(7)
                    remaining_run.font.color.rgb = RGBColor(0, 98, 110) if using_template else RGBColor(226, 232, 240)

    if roadmap_df is not None and not roadmap_df.empty:
        slide = prs.slides.add_slide(content_layout)
        add_chrome(slide, roadmap_title, roadmap_subtitle)
        earliest = roadmap_df["Milestone date"].min().replace(day=1)
        latest = roadmap_df["Milestone date"].max().replace(day=1) + pd.DateOffset(months=1)
        total_days = max((latest - earliest).days, 1)
        axis_left, axis_right, axis_y = 0.75, 12.55, 4.55

        def x_for_date(value):
            return axis_left + ((value - earliest).days / total_days) * (axis_right - axis_left)

        axis = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(axis_left), Inches(axis_y), Inches(axis_right - axis_left), Inches(0.025))
        axis.fill.solid()
        axis.fill.fore_color.rgb = white if using_template else muted
        axis.line.fill.background()
        month = earliest
        while month <= latest:
            x_pos = x_for_date(month)
            grid_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x_pos), Inches(3.95), Inches(0.012), Inches(0.68))
            grid_line.fill.solid()
            grid_line.fill.fore_color.rgb = template_teal if using_template else grid
            grid_line.line.fill.background()
            add_text(slide, month.strftime("%b %Y"), x_pos - 0.52, 3.58, 1.04, 0.20, size=11 if using_template else 8, color=template_muted if using_template else muted, bold=True, align=PP_ALIGN.CENTER)
            month += pd.DateOffset(months=1)

        today_x = x_for_date(pd.Timestamp.today())
        if axis_left <= today_x <= axis_right:
            today_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(today_x), Inches(2.0), Inches(0.018), Inches(2.65))
            today_line.fill.solid()
            today_line.fill.fore_color.rgb = RGBColor(239, 68, 68)
            today_line.line.fill.background()
            add_text(slide, "TODAY", today_x - 0.27, 1.83, 0.56, 0.16, size=7, color=RGBColor(239, 68, 68), bold=True, align=PP_ALIGN.CENTER)

        # One release card per milestone mirrors the supplied Release Calendar:
        # release version first, followed by the Epics planned for that release.
        displayed = roadmap_df.head(12).copy()
        displayed["Milestone date"] = pd.to_datetime(displayed["Milestone date"])
        release_cards = []
        for (milestone, release), rows in displayed.groupby(["Milestone date", "Release"], sort=True):
            release_cards.append({
                "date": milestone,
                "release": str(release or "Release TBD"),
                "epics": [
                    f"{str(key).strip()} - {str(summary).strip()}" if str(key).strip() else str(summary).strip()
                    for key, summary in zip(rows.get("Key", pd.Series("", index=rows.index)), rows["Summary"])
                    if str(summary).strip()
                ],
            })

        for index, card in enumerate(release_cards):
            marker_x = x_for_date(card["date"])
            box_width = 2.45
            box_height = min(1.25, 0.42 + (0.23 * min(len(card["epics"]), 3)))
            box_left = max(0.55, min(12.70 - box_width, marker_x - (box_width / 2)))
            box_top = 2.05 if index % 2 else 2.88
            card_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(box_left), Inches(box_top), Inches(box_width), Inches(box_height))
            card_shape.fill.solid()
            card_shape.fill.fore_color.rgb = template_background if using_template else RGBColor(255, 255, 255)
            card_shape.line.color.rgb = template_teal if using_template else primary
            add_text(slide, card["release"], box_left + 0.13, box_top + 0.08, box_width - 0.25, 0.18, size=9, color=template_teal if using_template else primary, bold=True)
            epic_lines = "\n".join(f"•  {epic}" for epic in card["epics"][:3])
            if len(card["epics"]) > 3:
                epic_lines += f"\n•  +{len(card['epics']) - 3} more"
            add_text(slide, epic_lines, box_left + 0.13, box_top + 0.27, box_width - 0.25, box_height - 0.32, size=7.7, color=white if using_template else dark)
            connector = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(marker_x), Inches(box_top + box_height), Inches(0.018), Inches(max(0.02, axis_y - (box_top + box_height))))
            connector.fill.solid()
            connector.fill.fore_color.rgb = template_teal if using_template else primary
            connector.line.fill.background()
            marker = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(marker_x - 0.09), Inches(axis_y - 0.09), Inches(0.18), Inches(0.18))
            marker.fill.solid()
            marker.fill.fore_color.rgb = template_teal if using_template else primary
            marker.line.fill.background()
        marker = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(0.70), Inches(6.42), Inches(0.14), Inches(0.14))
        marker.fill.solid()
        marker.fill.fore_color.rgb = template_teal if using_template else primary
        marker.line.fill.background()
        add_text(slide, "Release milestone", 0.92, 6.39, 1.5, 0.20, size=7.5, color=white if using_template else dark, bold=True)
        add_text(slide, f"Generated {pd.Timestamp.today().strftime('%d %b %Y')} - delivery dates are approximate", 0.48, 6.95, 5.5, 0.16, size=7.5, color=template_muted if using_template else muted)

    output = io.BytesIO()
    prs.save(output)
    output.seek(0)
    return output.getvalue()


def build_quarterly_plan_pdf(df, primary_color_hex):
    import io
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors

    pdf_buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        pdf_buffer,
        pagesize=landscape(letter),
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=22,
        textColor=colors.HexColor(primary_color_hex),
        spaceAfter=12
    )
    
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        textColor=colors.HexColor("#64748B"),
        spaceAfter=20
    )
    
    cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=11
    )
    
    cell_bold_style = ParagraphStyle(
        'TableCellBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=11
    )

    story = []
    story.append(Paragraph("🎯 QUARTERLY PLAN ROADMAP", title_style))
    proj_name = os.getenv("PROJECT_NAME", "RECALL2")
    today_str = pd.Timestamp.today().strftime('%d-%b-%Y')
    story.append(Paragraph(f"<b>Project:</b> {proj_name} | <b>Generated:</b> {today_str} | <b>Total Epics:</b> {len(df)}", subtitle_style))
    
    headers = ["Key", "Epic Name", "Status", "Sprint", "Size", "Quarter", "Cluster"]
    table_data = [[Paragraph(f"<b>{h}</b>", cell_bold_style) for h in headers]]
    
    sorted_df = df.copy()
    if 'Quarter' in sorted_df.columns:
        sorted_df = sorted_df.sort_values(by=['Quarter', 'Cluster', 'Sprint'], na_position='last')
        
    for _, row in sorted_df.iterrows():
        sprint_val = row.get("Sprint")
        sprint_str = f"Sprint {int(float(sprint_val))}" if pd.notna(sprint_val) and str(sprint_val).strip() != "" and str(sprint_val).strip().lower() != "nan" else "-"
        
        row_cells = [
            Paragraph(str(row.get("Key", "-")), cell_bold_style),
            Paragraph(str(row.get("Epic Name", "-")), cell_style),
            Paragraph(str(row.get("Status", "-")), cell_style),
            Paragraph(sprint_str, cell_style),
            Paragraph(str(row.get("Size", "-")), cell_style),
            Paragraph(str(row.get("Quarter", "-")), cell_style),
            Paragraph(str(row.get("Cluster", "-")), cell_style),
        ]
        table_data.append(row_cells)
        
    col_widths = [60, 260, 80, 80, 50, 90, 100]
    t = Table(table_data, colWidths=col_widths, repeatRows=1)
    
    t_style = TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor(primary_color_hex)),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ('TOPPADDING', (0,0), (-1,-1), 6),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.HexColor("#FFFFFF"), colors.HexColor("#F8FAFC")]),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
    ])
    
    for i in range(len(headers)):
        table_data[0][i].style.textColor = colors.white
        
    t.setStyle(t_style)
    story.append(t)
    
    doc.build(story)
    pdf_buffer.seek(0)
    return pdf_buffer


def get_sprint_duration():
    return st.session_state.get("sprint_duration", 15)

def get_size_mapping():
    return {
        'S': int(st.session_state.get("size_s_sprints", 1)),
        'M': int(st.session_state.get("size_m_sprints", 2)),
        'L': int(st.session_state.get("size_l_sprints", 3)),
        'XL': int(st.session_state.get("size_xl_sprints", 4))
    }

def calculate_sequential_dates(row):
    is_milestone = row.get('Milestone', False)
    
    epic_size = row.get('Size', 'M')
    size_mapping = get_size_mapping()
    sprint_duration = get_sprint_duration()
    
    if pd.isna(epic_size) or not isinstance(epic_size, str) or epic_size.upper() not in size_mapping:
        epic_size = 'M'
    else:
        epic_size = epic_size.upper()
        
    num_sprints = size_mapping[epic_size]
    
    # If it is a Milestone, the default duration is 1 day, otherwise calculated in sprints
    duration = 1 if is_milestone else (num_sprints * sprint_duration)
    
    sprint_val = row.get('Sprint')
    q_field = row.get('Quarter')
    base_date = parse_quarter_to_date(q_field)
    
    try:
        if pd.notna(sprint_val) and 'sprint_calendar' in st.session_state:
            sprint_num = int(float(sprint_val))
            if is_milestone:
                cal_row = st.session_state.sprint_calendar[st.session_state.sprint_calendar['Sprint'] == sprint_num]
                if not cal_row.empty:
                    start_date = pd.to_datetime(cal_row.iloc[0]['Start Date'])
                    due_date = start_date + pd.Timedelta(days=1)
                else:
                    start_date = base_date + pd.Timedelta(days=(sprint_num - 1) * sprint_duration)
                    due_date = start_date + pd.Timedelta(days=1)
            else:
                end_sprint_num = sprint_num + num_sprints - 1
                
                start_cal_row = st.session_state.sprint_calendar[st.session_state.sprint_calendar['Sprint'] == sprint_num]
                end_cal_row = st.session_state.sprint_calendar[st.session_state.sprint_calendar['Sprint'] == end_sprint_num]
                
                if not start_cal_row.empty:
                    start_date = pd.to_datetime(start_cal_row.iloc[0]['Start Date'])
                    if not end_cal_row.empty:
                        due_date = pd.to_datetime(end_cal_row.iloc[0]['Due Date'])
                    else:
                        due_date = start_date + pd.Timedelta(days=duration)
                else:
                    start_date = base_date + pd.Timedelta(days=(sprint_num - 1) * sprint_duration)
                    due_date = start_date + pd.Timedelta(days=duration)
        else:
            start_date = base_date
            due_date = start_date + pd.Timedelta(days=duration)
    except Exception as e:
        start_date = base_date
        due_date = start_date + pd.Timedelta(days=duration)
        
    return pd.Series([start_date, due_date], index=['Start Date', 'Due Date'])

def sync_sprint_calendar_changes():
    edits = st.session_state.get("sprint_cal_editor")
    if not edits or st.session_state.get("sprint_calendar") is None:
        return
    df = st.session_state.sprint_calendar.copy()
    for idx_str, col_changes in edits.get("edited_rows", {}).items():
        idx = int(idx_str)
        if idx < len(df):
            for col, val in col_changes.items():
                df.at[idx, col] = val
    st.session_state.sprint_calendar = df
    try:
        df.to_csv(SHARED_SPRINT_CALENDAR_PATH, index=False)
        st.session_state.last_sprint_cal_mtime = os.path.getmtime(SHARED_SPRINT_CALENDAR_PATH)
    except:
        pass
    
    # Recalculate main_df if present
    if 'main_df' in st.session_state:
        main_df = st.session_state.main_df.copy()
        for idx, row in main_df.iterrows():
            sprint_val = row.get('Sprint')
            if pd.notna(sprint_val):
                try:
                    s_num = int(float(sprint_val))
                    cal_row = df[df['Sprint'] == s_num]
                    if not cal_row.empty:
                        calculated = calculate_sequential_dates(row)
                        main_df.at[idx, 'Start Date'] = calculated[0]
                        main_df.at[idx, 'Due Date'] = calculated[1]
                except:
                    pass
        st.session_state.main_df = main_df
        try:
            main_df.to_csv(SHARED_BACKLOG_PATH, index=False)
            st.session_state.last_backlog_mtime = os.path.getmtime(SHARED_BACKLOG_PATH)
        except:
            pass


def sync_main_editor_changes():
    editor_key = f"main_editor_{st.session_state.get('editor_key_counter', 0)}"
    edits = st.session_state.get(editor_key)
    if not edits or st.session_state.get("main_df") is None:
        return
    
    df = st.session_state.main_df.copy()
    sprint_duration = get_sprint_duration()
    size_mapping = get_size_mapping()
    
    # Process Edited Rows
    for idx_str, col_changes in edits.get("edited_rows", {}).items():
        idx = int(idx_str)
        if idx in df.index:
            old_row = df.loc[idx]
            new_row = old_row.copy()
            for col, val in col_changes.items():
                new_row[col] = val
                
            sprint_changed = val_changed(old_row.get('Sprint'), new_row.get('Sprint'))
            quarter_changed = val_changed(old_row.get('Quarter'), new_row.get('Quarter'))
            milestone_changed = val_changed(old_row.get('Milestone'), new_row.get('Milestone'))
            size_changed = val_changed(old_row.get('Size'), new_row.get('Size'))
            start_date_changed = val_changed(old_row.get('Start Date'), new_row.get('Start Date'))
            due_date_changed = val_changed(old_row.get('Due Date'), new_row.get('Due Date'))
            dates_null = bool(pd.isna(new_row.get('Start Date')) or pd.isna(new_row.get('Due Date')))
            
            if sprint_changed or quarter_changed or milestone_changed or size_changed or dates_null:
                calculated = calculate_sequential_dates(new_row)
                new_row['Start Date'] = calculated[0]
                new_row['Due Date'] = calculated[1]
            elif start_date_changed or due_date_changed:
                new_start = new_row.get('Start Date')
                new_due = new_row.get('Due Date')
                
                if pd.notna(new_start):
                    new_start_dt = pd.to_datetime(new_start)
                    matching_sprint = pd.NA
                    if 'sprint_calendar' in st.session_state:
                        for _, cal_row in st.session_state.sprint_calendar.iterrows():
                            cal_start = pd.to_datetime(cal_row['Start Date'])
                            cal_due = pd.to_datetime(cal_row['Due Date'])
                            if cal_start <= new_start_dt < cal_due:
                                matching_sprint = int(cal_row['Sprint'])
                                break
                    new_row['Sprint'] = matching_sprint
                    
                if pd.notna(new_start) and pd.notna(new_due):
                    new_start_dt = pd.to_datetime(new_start)
                    new_due_dt = pd.to_datetime(new_due)
                    duration_days = (new_due_dt - new_start_dt).days
                    approx_sprints = max(1, round(duration_days / sprint_duration))
                    
                    best_size = 'M'
                    min_diff = 999
                    for sz, sz_sprints in size_mapping.items():
                        diff = abs(sz_sprints - approx_sprints)
                        if diff < min_diff:
                            min_diff = diff
                            best_size = sz
                    new_row['Size'] = best_size
            
            for col in df.columns:
                df.at[idx, col] = new_row[col]

    # Process Deleted Rows
    deleted_indices = edits.get("deleted_rows", [])
    if deleted_indices:
        df.drop(index=deleted_indices, inplace=True)
        df.reset_index(drop=True, inplace=True)

    # Process Added Rows
    added_rows = edits.get("added_rows", [])
    if added_rows:
        new_rows_list = []
        for added_row in added_rows:
            row_dict = {
                'Key': added_row.get('Key', 'KEY-999'),
                'Epic Name': added_row.get('Epic Name', 'New Epic'),
                'Status': added_row.get('Status', 'To Do'),
                'Sprint': added_row.get('Sprint', pd.NA),
                'Size': added_row.get('Size', 'M'),
                'Quarter': added_row.get('Quarter', 'Q1 - 26'),
                'Cluster': added_row.get('Cluster', 'General'),
                'Cluster Name': added_row.get('Cluster Name', ''),
                'Milestone': added_row.get('Milestone', False),
                'Labels': added_row.get('Labels', ''),
                'Description': added_row.get('Description', 'No description provided.')
            }
            for k, v in added_row.items():
                if k not in row_dict:
                    row_dict[k] = v
                    
            calculated = calculate_sequential_dates(row_dict)
            row_dict['Start Date'] = calculated[0]
            row_dict['Due Date'] = calculated[1]
            new_rows_list.append(row_dict)
            
        new_rows_df = pd.DataFrame(new_rows_list)
        for col in df.columns:
            if col not in new_rows_df.columns:
                new_rows_df[col] = pd.NA
        new_rows_df = new_rows_df[df.columns]
        df = pd.concat([df, new_rows_df], ignore_index=True)

    st.session_state.main_df = df
    try:
        df.to_csv(SHARED_BACKLOG_PATH, index=False)
        st.session_state.last_backlog_mtime = os.path.getmtime(SHARED_BACKLOG_PATH)
    except:
        pass

st.title("🎯 Quarterly Planner")

# Load credentials and status
jira_status = st.session_state.get("jira_connection_status")
jira_msg = st.session_state.get("jira_connection_msg", "Not checked yet")
jira_server = st.session_state.get("jira_server", "")

# ---------------------------------------------------------
# Data Loader and Normalizer Helper
# ---------------------------------------------------------
def load_and_normalize_data(file_or_path):
    df_raw = pd.read_csv(file_or_path)
    
    # Drop rows that are completely empty
    if not df_raw.empty:
        df_raw.dropna(how='all', inplace=True)
        
    # Normalize common column headers from Jira CSV exports
    rename_mapping = {}
    for col in df_raw.columns:
        col_lower = col.strip().lower()
        if col_lower in ['issue key', 'issue_key', 'key', 'id']:
            rename_mapping[col] = 'Key'
        elif col_lower in ['summary', 'epic name', 'epic_name', 'title']:
            rename_mapping[col] = 'Epic Name'
        elif col_lower in ['status', 'status name', 'state']:
            rename_mapping[col] = 'Status'
        elif col_lower in ['description', 'desc', 'detailed description']:
            rename_mapping[col] = 'Description'
        elif col_lower in ['sprint', 'sprint number']:
            rename_mapping[col] = 'Sprint'
        elif col_lower in ['size', 'epic size', 'epic_size', 'talla', 'estimacion', 'estimación']:
            rename_mapping[col] = 'Size'
        elif col_lower in ['quarter', 'period']:
            rename_mapping[col] = 'Quarter'
        elif col_lower in ['cluster name', 'cluster_name', 'cluster-name']:
            rename_mapping[col] = 'Cluster Name'
        elif col_lower in ['cluster', 'stream', 'area']:
            rename_mapping[col] = 'Cluster'
        elif col_lower in ['labels', 'label', 'tags']:
            rename_mapping[col] = 'Labels'
            
    if rename_mapping:
        df_raw.rename(columns=rename_mapping, inplace=True)
        
    # Ensure minimum columns and clean missing values (NaN) to prevent crashes
    expected_cols = {
        'Key': 'KEY-999',
        'Epic Name': 'New Epic',
        'Status': 'To Do',
        'Sprint': pd.NA,
        'Size': 'M',
        'Quarter': 'Q1 - 26',
        'Cluster': 'General',
        'Labels': '',
        'Description': 'No description provided.',
        'Milestone': False
    }
    for col, default_val in expected_cols.items():
        if col not in df_raw.columns:
            df_raw[col] = default_val
        else:
            if col == 'Milestone':
                df_raw[col] = df_raw[col].fillna(False).astype(bool)
            elif col == 'Sprint':
                df_raw[col] = pd.to_numeric(df_raw[col], errors='coerce')
            elif col == 'Size':
                df_raw[col] = df_raw[col].fillna('M').astype(str).str.strip().str.upper()
                df_raw[col] = df_raw[col].apply(lambda x: x if x in ['S', 'M', 'L', 'XL'] else 'M')
            elif col in ['Key', 'Epic Name', 'Status', 'Quarter', 'Cluster', 'Description', 'Labels']:
                df_raw[col] = df_raw[col].fillna(default_val)
                
    if 'Cluster Name' in df_raw.columns:
        df_raw['Cluster Name'] = df_raw['Cluster Name'].fillna("")
        
    # Fill missing dates automatically
    if 'Start Date' not in df_raw.columns:
        df_raw['Start Date'] = pd.NaT
    if 'Due Date' not in df_raw.columns:
        df_raw['Due Date'] = pd.NaT
        
    df_raw['Start Date'] = pd.to_datetime(df_raw['Start Date'], errors='coerce')
    df_raw['Due Date'] = pd.to_datetime(df_raw['Due Date'], errors='coerce')
    
    mask_missing = df_raw['Start Date'].isna() | df_raw['Due Date'].isna()
    if mask_missing.any():
        calculated_dates = df_raw[mask_missing].apply(calculate_sequential_dates, axis=1, result_type='expand')
        if not calculated_dates.empty:
            df_raw.loc[mask_missing, ['Start Date', 'Due Date']] = calculated_dates
            
    return df_raw

# ---------------------------------------------------------
# Sidebar Navigation Steps
# ---------------------------------------------------------
options_qp = ["🔌 Ingestion", "✍️ Workbook", "💾 Exporter"]

if 'active_tab_qp' not in st.session_state or st.session_state.active_tab_qp not in options_qp:
    st.session_state.active_tab_qp = "🔌 Ingestion"

selected_nav = st.sidebar.radio(
    "Select current step:",
    options=options_qp,
    index=options_qp.index(st.session_state.active_tab_qp)
)

if selected_nav != st.session_state.active_tab_qp:
    st.session_state.active_tab_qp = selected_nav
    st.rerun()

# Auto-load shared backlog on startup if not already loaded
if 'main_df' not in st.session_state:
    if os.path.exists(SHARED_BACKLOG_PATH):
        try:
            st.session_state.main_df = load_and_normalize_data(SHARED_BACKLOG_PATH)
            st.session_state.last_backlog_mtime = os.path.getmtime(SHARED_BACKLOG_PATH)
        except:
            pass

# ---------------------------------------------------------
# STEP 1: Ingestion
# ---------------------------------------------------------
if st.session_state.active_tab_qp == "🔌 Ingestion":
    st.subheader("🔌 Dual Jira & CSV Backlog Ingestion")
    st.write("Upload your product backlog CSV or connect to your Jira server parameters to get started.")
    
    # Ingestion Flash Feedback
    if "ingestion_feedback" in st.session_state:
        fb = st.session_state.ingestion_feedback
        if fb.get("type") == "success":
            st.success(fb.get("text"))
        elif fb.get("type") == "error":
            st.error(fb.get("text"))
        elif fb.get("type") == "warning":
            st.warning(fb.get("text"))
        del st.session_state.ingestion_feedback
    
    # Connection status display
    if not jira_server:
        st.warning("⚠️ **Jira Connection**: Not configured. Set credentials in **Home Hub** -> **Centralized Integrations** tab.")
    elif jira_status == "Success":
        st.success(f"🟢 **Jira Connected**: `{jira_server}` ({jira_msg})")
    elif jira_status == "Failed":
        st.error(f"🔴 **Jira Connection Failed**: `{jira_server}` ({jira_msg}). Fix it on the **Home Hub**.")
    else:
        st.info(f"🟡 **Jira Configured**: `{jira_server}` (Status: {jira_msg}). Go to **Home Hub** to test connection.")
        
    st.divider()
    
    uploaded_file = st.file_uploader("Upload your data (CSV)", type=['csv'])
    if uploaded_file:
        try:
            st.session_state.main_df = load_and_normalize_data(uploaded_file)
            st.session_state.main_df.to_csv(SHARED_BACKLOG_PATH, index=False)
            st.session_state.last_backlog_mtime = os.path.getmtime(SHARED_BACKLOG_PATH)
            st.session_state.editor_key_counter = st.session_state.get('editor_key_counter', 0) + 1
            st.session_state.ingestion_feedback = {"type": "success", "text": f"✅ Backlog of {len(st.session_state.main_df)} rows loaded and normalized successfully!"}
            st.rerun()
        except Exception as e:
            st.session_state.ingestion_feedback = {"type": "error", "text": f"❌ Failed to parse CSV file: {e}"}
            st.rerun()
            
    # Success indicator if data is loaded
    if st.session_state.get("main_df") is not None:
        st.info(f"📊 **Current Loaded Backlog:** `{len(st.session_state.main_df)}` rows found in active workspace.")
        
    st.markdown("### Don't have a file at hand?")
    st.write("We have created a realistic mock file with simulated data so you can test the application with one click:")
    try:
        mock_path_parent = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "jira_mock.csv")
        mock_path_tools = os.path.join(os.path.dirname(os.path.abspath(__file__)), "jira_mock.csv")
        mock_path = mock_path_parent if os.path.exists(mock_path_parent) else mock_path_tools
        with open(mock_path, "r") as f:
            mock_csv = f.read()
        st.download_button(
            label="⬇️ Download a Template", 
            data=mock_csv.encode('utf-8'), 
            file_name="jira_mock.csv", 
            mime="text/csv"
        )
    except:
        pass

    st.divider()
    st.subheader("📈 Quarterly Epic Progress")
    st.write("Load the committed Epics for the quarter, calculate their completion, and add a concise update for the presentation.")
    with st.container(border=True):
        col_project, col_committed, col_quarter = st.columns(3)
        with col_project:
            progress_project = st.text_input("Project", value="RECALLTWO", key="qp_progress_project")
        with col_committed:
            progress_committed_label = st.text_input(
                "Committed Epic label", value=os.getenv("COMMITTED_LABEL", "RC2_committed"), key="qp_progress_committed"
            )
        with col_quarter:
            progress_quarter_label = st.text_input(
                "Quarter label", value=os.getenv("QUARTER_LABEL", "RC2_FB_18"), key="qp_progress_quarter"
            )
        progress_title = st.text_input("Progress slide title", value="Committed Epics Q3 - Status", key="qp_progress_title")
        progress_subtitle = st.text_input("Progress slide subtitle", value="Detail overview of Epics Progress for Q3", key="qp_progress_subtitle")
        roadmap_title = st.text_input("Roadmap slide title", value="Delivery Roadmap", key="qp_roadmap_title")
        roadmap_subtitle = st.text_input("Roadmap slide subtitle", value="Planned delivery milestones for unfinished work", key="qp_roadmap_subtitle")

        if st.button("📈 Load Quarterly Epic Progress", use_container_width=True):
            if not st.session_state.get("jira_server") or not st.session_state.get("jira_token"):
                st.error("Configure and test the Jira connection first in Home Hub.")
            elif not progress_committed_label.strip() or not progress_quarter_label.strip():
                st.error("Enter both the committed Epic label and the quarter label.")
            else:
                with st.spinner("Downloading committed Epics and calculating progress..."):
                    result = build_quarterly_epic_progress_table(
                        st.session_state.jira_server,
                        st.session_state.jira_token,
                        progress_committed_label.strip(),
                        progress_quarter_label.strip(),
                        progress_title.strip() or "Quarterly Epic Progress",
                        "Quarterly Planner",
                        st.session_state.get("jira_auth_method", "Personal Access Token (Bearer PAT)"),
                        st.session_state.get("jira_email", ""),
                        project_key=progress_project.strip() or "RECALLTWO",
                    )
                if result is None:
                    st.error("The quarterly Epic progress could not be loaded. Check the Jira connection and labels.")
                elif result["df"].empty:
                    st.warning("No committed Epics were found for those labels.")
                else:
                    previous = st.session_state.get("quarterly_progress_df")
                    previous_updates = {}
                    if isinstance(previous, pd.DataFrame) and not previous.empty and "Presentation update" in previous.columns:
                        previous_updates = previous.set_index("Key")["Presentation update"].fillna("").to_dict()
                    progress_df = result["df"].drop(columns=["Select"], errors="ignore")
                    progress_df["Presentation update"] = progress_df["Key"].map(previous_updates).fillna("")
                    st.session_state.quarterly_progress_df = prepare_quarterly_progress_for_presentation(progress_df)
                    st.session_state.quarterly_progress_config = {
                        "title": progress_title.strip() or "Quarterly Epic Progress",
                        "subtitle": progress_subtitle.strip(),
                        "roadmap_title": roadmap_title.strip() or "Delivery Roadmap",
                        "roadmap_subtitle": roadmap_subtitle.strip(),
                        "quarter_label": progress_quarter_label.strip(),
                    }
                    st.success(f"Loaded {len(progress_df)} committed Epics. Add the presentation updates in Workbook.")
                    st.rerun()





# ---------------------------------------------------------
# STEP 2: Workbook
# ---------------------------------------------------------
elif st.session_state.active_tab_qp == "✍️ Workbook":
    st.subheader("✍️ Backlog Workbook Editor")
    st.write("Modify your epics, status, sprints, and period schedules directly in the high-fidelity table below.")

    progress_df = st.session_state.get("quarterly_progress_df")
    st.markdown("### 📈 Quarterly Epic Progress — presentation updates")
    if isinstance(progress_df, pd.DataFrame) and not progress_df.empty:
        st.caption("Select the Epics to communicate, define their planned release milestone, and write a short update. Jira fields remain read-only.")
        progress_columns = [
            "Key", "Summary", "Completion", "Closed / Resolved", "To Do", "In progress", "Status", "Team", "Presentation update",
            "Include in delivery roadmap", "Release version", "Delivery month", "Delivery timing"
        ]
        progress_columns = [column for column in progress_columns if column in progress_df.columns]
        with st.form("quarterly_progress_updates_form", border=False):
            edited_progress = st.data_editor(
                progress_df[progress_columns],
                use_container_width=True,
                hide_index=True,
                disabled=[column for column in progress_columns if column not in {
                    "Presentation update", "Include in delivery roadmap", "Release version", "Delivery month", "Delivery timing"
                }],
                column_config={
                    "Key": st.column_config.TextColumn("Key", width="small"),
                    "Summary": st.column_config.TextColumn("Epic", width="large"),
                "Completion": st.column_config.TextColumn("Completion", width="small"),
                "Closed / Resolved": st.column_config.NumberColumn("Closed / Resolved", width="small", format="%d"),
                "To Do": st.column_config.NumberColumn("To Do", width="small", format="%d"),
                "In progress": st.column_config.NumberColumn("In progress", width="small", format="%d"),
                    "Status": st.column_config.TextColumn("Status", width="small"),
                    "Team": st.column_config.TextColumn("Team", width="medium"),
                    "Presentation update": st.column_config.TextColumn("Update / context", width="large"),
                    "Include in delivery roadmap": st.column_config.CheckboxColumn("Show in roadmap", default=False),
                    "Release version": st.column_config.TextColumn("Release version", width="medium"),
                    "Delivery month": st.column_config.SelectboxColumn("Delivery month", options=delivery_month_options(), width="medium"),
                    "Delivery timing": st.column_config.SelectboxColumn("Approx. timing", options=["Beginning", "Mid", "End"], width="small"),
                },
                key="quarterly_progress_editor",
            )
            save_progress_updates = st.form_submit_button("Save roadmap selections and updates", use_container_width=True)
        if save_progress_updates:
            updated_progress = progress_df.copy()
            for column in ["Presentation update", "Include in delivery roadmap", "Release version", "Delivery month", "Delivery timing"]:
                updated_progress[column] = edited_progress[column]
            st.session_state.quarterly_progress_df = prepare_quarterly_progress_for_presentation(updated_progress)
            st.success("Roadmap selections and presentation updates saved.")
    else:
        st.info("Load Quarterly Epic Progress in Ingestion to add the updates that will appear in the slide.")
    st.divider()
    
    if st.session_state.get("main_df") is None:
        st.warning("⚠️ **No backlog data loaded**. Please complete Step 1: **Ingestion** or load a template backlog first.")
    else:
        # Search & Quick Filters inside the workbook page
        st.write("### 🔍 Search Backlog")
        wb_search = st.text_input("Filter table by text:", value="", placeholder="Type summary key or cluster...")
        
        df_editor_source = st.session_state.main_df.copy()
        if wb_search:
            q = wb_search.lower()
            mask = (
                df_editor_source['Epic Name'].astype(str).str.lower().str.contains(q) |
                df_editor_source['Key'].astype(str).str.lower().str.contains(q) |
                df_editor_source['Description'].astype(str).str.lower().str.contains(q) |
                df_editor_source['Cluster'].astype(str).str.lower().str.contains(q)
            )
            df_editor_source = df_editor_source[mask]

        column_config = {
            "Key": st.column_config.TextColumn("id", width="medium"),
            "Epic Name": st.column_config.TextColumn("Epic Name", width="large"),
            "Status": st.column_config.SelectboxColumn("Status", options=["To Do", "In Progress", "Done", "Blocked", "Open", "Closed"]),
            "Sprint": st.column_config.NumberColumn("Sprint", min_value=1, max_value=20, step=1),
            "Size": st.column_config.SelectboxColumn("Size", options=["S", "M", "L", "XL"], default="M", width="small"),
            "Quarter": st.column_config.TextColumn("Quarter"),
            "Cluster": st.column_config.TextColumn("Cluster"),
            "Cluster Name": st.column_config.TextColumn("Cluster Name"),
            "Milestone": st.column_config.CheckboxColumn("Milestone", default=False),
            "Start Date": st.column_config.DatetimeColumn("Start Date", format="DD-MM-YYYY"),
            "Due Date": st.column_config.DatetimeColumn("Due Date", format="DD-MM-YYYY")
        }
        
        st.data_editor(
            df_editor_source, 
            num_rows="dynamic", 
            use_container_width=True,
            column_config=column_config,
            key=f"main_editor_{st.session_state.get('editor_key_counter', 0)}",
            on_change=sync_main_editor_changes
        )
        
        st.info("💡 **Tip:** Changes are saved automatically to the shared backlog workspace files in real time.")
        
        # Action: Recalculate all dates button
        st.markdown("---")
        st.markdown("#### 🔄 Recalculate Planning Dates")
        st.write("Recalculate start and end planning dates for all backlog items based on their Sizes, Sprint calendars, and Quarter specifications:")
        if st.button("Recalculate All Dates", use_container_width=True):
            calculated = st.session_state.main_df.apply(calculate_sequential_dates, axis=1, result_type='expand')
            st.session_state.main_df[['Start Date', 'Due Date']] = calculated
            try:
                st.session_state.main_df.to_csv(SHARED_BACKLOG_PATH, index=False)
                st.session_state.last_backlog_mtime = os.path.getmtime(SHARED_BACKLOG_PATH)
            except:
                pass
def show_timeline_gantt_tab():
    df_for_filters = st.session_state.main_df
    
    # 2. Sidebar: Gantt Filters (AND)
    st.sidebar.header("🔍 Gantt Filters")
    search_query = st.sidebar.text_input("Search by Name or Key:", value="")
    
    # Dynamic filter by Quarter
    quarters_sel = []
    if 'Quarter' in df_for_filters.columns:
        all_quarters = sorted(df_for_filters['Quarter'].dropna().unique())
        quarters_sel = st.sidebar.multiselect("Filter by Quarter:", all_quarters)
        
    # Dynamic filter by Cluster
    clusters_sel = []
    if 'Cluster' in df_for_filters.columns:
        all_clusters = sorted(df_for_filters['Cluster'].dropna().unique())
        clusters_sel = st.sidebar.multiselect("Filter by Cluster:", all_clusters)

    # Dynamic filter by Labels
    labels_sel = []
    if 'Labels' in df_for_filters.columns:
        all_labels = set()
        for row in df_for_filters['Labels'].dropna(): 
            all_labels.update(str(row).split())
        labels_sel = st.sidebar.multiselect("Filter by Labels:", sorted(list(all_labels)))

    # Filter by Status
    status_sel = []
    if 'Status' in df_for_filters.columns:
        all_statuses = sorted(df_for_filters['Status'].dropna().unique())
        status_sel = st.sidebar.multiselect("Filter by Status:", all_statuses)

    # Sidebar: Collapsible Settings
    st.sidebar.divider()
    st.sidebar.header("⚙️ Settings")

    with st.sidebar.expander("⚙️ Sprint Duration", expanded=False):
        sprint_duration = st.number_input("Days:", min_value=1, max_value=90, value=15, key="sprint_duration")
        
    with st.sidebar.expander("📐 Sizes to Sprints Equivalence", expanded=False):
        size_s_sprints = st.number_input("S Size:", min_value=1, max_value=20, value=1, key="size_s_sprints")
        size_m_sprints = st.number_input("M Size:", min_value=1, max_value=20, value=2, key="size_m_sprints")
        size_l_sprints = st.number_input("L Size:", min_value=1, max_value=20, value=3, key="size_l_sprints")
        size_xl_sprints = st.number_input("XL Size:", min_value=1, max_value=20, value=4, key="size_xl_sprints")
        
    size_mapping = get_size_mapping()

    # Initialize the Sprint Calendar in Session State if it doesn't exist
    if 'sprint_calendar' not in st.session_state:
        q_sample = "Q1 - 26"
        non_null_q = df_for_filters['Quarter'].dropna()
        if not non_null_q.empty:
            q_sample = non_null_q.iloc[0]
            
        base_date = parse_quarter_to_date(q_sample)
        
        sprint_data = []
        for i in range(1, 7):
            s_start = base_date + pd.Timedelta(days=(i - 1) * sprint_duration)
            s_end = s_start + pd.Timedelta(days=sprint_duration)
            sprint_data.append({
                "Sprint": int(i),
                "Start Date": s_start.date(),
                "Due Date": s_end.date()
            })
        st.session_state.sprint_calendar = pd.DataFrame(sprint_data)
        
    # Reactivity of sprint duration over the calendar
    if 'prev_sprint_duration' not in st.session_state or st.session_state.prev_sprint_duration != sprint_duration:
        st.session_state.prev_sprint_duration = sprint_duration
        base_date = pd.to_datetime(st.session_state.sprint_calendar.iloc[0]['Start Date'])
        updated_sprints = []
        for idx, row in st.session_state.sprint_calendar.iterrows():
            s_num = int(row['Sprint'])
            s_start = base_date + pd.Timedelta(days=(s_num - 1) * sprint_duration)
            s_end = s_start + pd.Timedelta(days=sprint_duration)
            updated_sprints.append({
                "Sprint": s_num,
                "Start Date": s_start.date(),
                "Due Date": s_end.date()
            })
        st.session_state.sprint_calendar = pd.DataFrame(updated_sprints)

    with st.sidebar.expander("📅 Sprints Dates", expanded=False):
        st.data_editor(
            st.session_state.sprint_calendar,
            num_rows="fixed",
            use_container_width=True,
            column_config={
                "Sprint": st.column_config.NumberColumn("Sprint", disabled=True),
                "Start Date": st.column_config.DateColumn("Start Date"),
                "Due Date": st.column_config.DateColumn("End Date")
            },
            key="sprint_cal_editor",
            on_change=sync_sprint_calendar_changes
        )

    # Apply filters mask to a copy for visual display and analysis
    filtered_df = st.session_state.main_df.copy()
    
    if search_query:
        query = search_query.lower()
        mask_name = filtered_df['Epic Name'].astype(str).str.lower().str.contains(query)
        mask_key = filtered_df['Key'].astype(str).str.lower().str.contains(query)
        mask_desc = filtered_df['Description'].astype(str).str.lower().str.contains(query)
        filtered_df = filtered_df[mask_name | mask_key | mask_desc]
        
    if quarters_sel:
        filtered_df = filtered_df[filtered_df['Quarter'].isin(quarters_sel)]
        
    if clusters_sel:
        filtered_df = filtered_df[filtered_df['Cluster'].isin(clusters_sel)]
        
    if labels_sel:
        mask_labels = filtered_df['Labels'].apply(lambda x: any(e in str(x).split() for e in labels_sel) if pd.notna(x) else False)
        filtered_df = filtered_df[mask_labels]

    if status_sel:
        filtered_df = filtered_df[filtered_df['Status'].isin(status_sel)]

    # Real-time Quality Validation Engine
    def validate_plan_data(df):
        errors = []
        warnings = []
        
        for idx, row in df.iterrows():
            epic_name = row.get('Epic Name', f"Row {idx + 1}")
            key = row.get('Key', f"Row {idx + 1}")
            label = f"[{key}] {epic_name}"
            
            start = row.get('Start Date')
            due = row.get('Due Date')
            
            # Validation 1: Due date before start date
            if pd.notna(start) and pd.notna(due):
                if due < start:
                    errors.append(f"❌ **{label}**: The due date ({due.strftime('%d-%m-%Y')}) is earlier than the start date ({start.strftime('%d-%m-%Y')}).")
            
            # Validation 2: Extreme Quarter deviation (supports natural multi-quarter span)
            quarter = row.get('Quarter')
            if pd.notna(quarter) and isinstance(quarter, str) and pd.notna(start) and pd.notna(due):
                base_date = parse_quarter_to_date(quarter)
                limit_start = base_date - pd.DateOffset(months=1)
                limit_end = base_date + pd.DateOffset(months=6)
                
                if start < limit_start or due > limit_end:
                    warnings.append(f"⚠️ **{label}**: The scheduled dates exceed the reasonable limit of quarter **{quarter}** (deviation greater than 2 quarters). Please verify if the planning is correct.")
            
            # Validation 3: Date out of assigned Sprint bounds
            sprint_val = row.get('Sprint')
            if pd.notna(sprint_val) and 'sprint_calendar' in st.session_state:
                try:
                    s_idx = int(float(sprint_val))
                    cal_row = st.session_state.sprint_calendar[st.session_state.sprint_calendar['Sprint'] == s_idx]
                    if not cal_row.empty:
                        s_start = pd.to_datetime(cal_row.iloc[0]['Start Date'])
                        s_due = pd.to_datetime(cal_row.iloc[0]['Due Date'])
                        if pd.notna(start) and (start < s_start or start >= s_due):
                            warnings.append(f"⚠️ **{label}**: Start date ({start.strftime('%d-%m-%Y')}) lies outside the bounds of assigned Sprint {s_idx} ({s_start.strftime('%d-%m-%Y')} to {s_due.strftime('%d-%m-%Y')}).")
                except:
                    pass
                    
        return errors, warnings

    critical_errors, control_warnings = validate_plan_data(st.session_state.main_df)
    
    # Render validation anomalies if any exist
    if critical_errors or control_warnings:
        total_alerts = len(critical_errors) + len(control_warnings)
        with st.expander(f"⚠️ **Quality & Date Validation Panel** ({total_alerts} Alerts)", expanded=False):
            st.markdown(
                '<div class="validation-card">'
                '<strong>Consistency Control:</strong> The validation engine has detected planning anomalies in real time. Please review them to ensure plan coherence.'
                '</div>',
                unsafe_allow_html=True
            )
            
            if critical_errors:
                st.error("🚨 **Critical Errors (Must be corrected in the table):**")
                for err in critical_errors:
                    st.markdown(f"- {err}")
                    
            if control_warnings:
                st.warning("⚠️ **Warnings of Misalignment / Limits:**")
                for adv in control_warnings:
                    st.markdown(f"- {adv}")

    
# ---------------------------------------------------------
    # 7. Gantt Chart View (Plotly)
    # ---------------------------------------------------------
    with st.expander("📅 Gantt Chart", expanded=True):
        if not filtered_df.empty:
            gantt_df = filtered_df.copy()
    
            # Standard colors mapped to roadmap statuses
            status_colors = {
                "To Do": "#94A3B8",       
                "In Progress": "#60A5FA",  
                "Done": "#4ADE80",        
                "Blocked": "#F87171",     
                "Milestone": "#FACC15",   
                "Cluster Header": "#475569" # Dark neutral header color for Cluster Summary rollup bar
            }
            
            gantt_df['Gantt_Status'] = gantt_df.apply(
                lambda r: get_status_group(r.get('Status', 'To Do'), r.get('Milestone') == True), axis=1
            )
            gantt_df['Sprint_Visual'] = gantt_df['Sprint'].apply(lambda x: f"Sprint {int(x)}" if pd.notna(x) else "Unassigned")
            
            # Inicializar estados de colapso si no existen
            if 'collapsed_clusters' not in st.session_state:
                st.session_state.collapsed_clusters = set()
            if 'chart_key_counter' not in st.session_state:
                st.session_state.chart_key_counter = 0
    
            unique_clusters = gantt_df['Cluster'].fillna("General").unique()
            sorted_clusters = sorted([c for c in unique_clusters if c != "General"])
            if "General" in unique_clusters:
                sorted_clusters.append("General")
    
            # Top-level UI controls for expanding/collapsing
            st.markdown('### 👀 Gantt View Mode')
            
            # Inject CSS specifically for the radio buttons to change their text color
            st.markdown(f"""
                <style>
                /* Apunta a las etiquetas del radio button */
                div.stRadio > div[role="radiogroup"] p {{
                    color: {st.session_state.get("primary_color", "#3B82F6")} !important; /* Accent */
                    font-weight: 600 !important;
                    font-size: 1.05rem !important;
                }}
                </style>
            """, unsafe_allow_html=True)
            
            view_mode = st.radio(
                "Selecciona el modo de visualización",
                options=["Clusters", "Items", "All"],
                index=2, # Default to "All"
                horizontal=True,
                label_visibility="collapsed"
            )
    
            # Detect active year and determine shared X-axis range
            años_detectados = gantt_df['Start Date'].dt.year.dropna().unique()
            año_activo = int(años_detectados[0]) if len(años_detectados) > 0 else pd.Timestamp.today().year
    
            shared_xaxis_range = None
            if quarters_sel:
                start_dates = []
                end_dates = []
                for q in quarters_sel:
                    q_date = parse_quarter_to_date(q)
                    start_dates.append(q_date)
                    end_dates.append(q_date + pd.DateOffset(months=3))
                if start_dates and end_dates:
                    shared_xaxis_range = [min(start_dates), max(end_dates)]
            else:
                shared_xaxis_range = [f"{año_activo}-01-01", f"{año_activo}-12-31"]
    
            plotting_rows = []
            y_axis_order = []
            epic_counter = 0
            
            # Build list rows to establish visual hierarchy in Gantt chart
            for idx_c, cluster_name in enumerate(sorted_clusters):
                cluster_gantt_df = gantt_df[gantt_df['Cluster'].fillna("General") == cluster_name]
                if cluster_gantt_df.empty: 
                    continue
                
                # Compute auto-rollup start and end dates for each Cluster
                c_start = cluster_gantt_df['Start Date'].min()
                c_due = cluster_gantt_df['Due Date'].max()
                if pd.isna(c_start): c_start = pd.Timestamp.today()
                if pd.isna(c_due): c_due = pd.Timestamp.today() + pd.DateOffset(days=30)
                
                cluster_name_full = cluster_name
                if 'Cluster Name' in cluster_gantt_df.columns:
                    non_null_names = cluster_gantt_df['Cluster Name'].dropna()
                    non_null_names = [n for n in non_null_names if str(n).strip() != ""]
                    if non_null_names:
                        cluster_name_full = f"{cluster_name} - {non_null_names[0]}"
                
                if view_mode in ["Clusters", "All"]:
                    num_items = len(cluster_gantt_df)
                    # Split cluster text into two lines: Cluster Key/Name on top, sub-name on bottom
                    cluster_desc = ""
                    if 'Cluster Name' in cluster_gantt_df.columns:
                        non_null_names = cluster_gantt_df['Cluster Name'].dropna()
                        non_null_names = [n for n in non_null_names if str(n).strip() != ""]
                        if non_null_names:
                            cluster_desc = f"<br><span style='font-size:10px; color:#94A3B8;'>{non_null_names[0]}</span>"
                    
                    cluster_y_label = f"📁 <b>{cluster_name.upper()}</b> ({num_items}){cluster_desc}"
                    cluster_id = f"C_{idx_c}"
                    
                    # Append the Cluster summary row
                    plotting_rows.append({
                        'Key': f"CLUSTER_{cluster_name}",
                        'Epic Name': f"Cluster: {cluster_name_full}",
                        'Status': 'N/A',
                        'Quarter': 'N/A',
                        'Sprint_Visual': 'N/A',
                        'Cluster': cluster_name,
                        'Start Date': c_start,
                        'Due Date': c_due,
                        'Gantt_Status': 'Cluster Header',
                        'Y_Axis_Label': cluster_id,
                        'Display_Label': cluster_y_label,
                        'Is_Cluster_Header': True,
                        'Y_Index': epic_counter,
                        'Size': 'N/A'
                    })
                    y_axis_order.append(cluster_id)
                    epic_counter += 1
                
                # If view mode includes items, inject the individual epics
                if view_mode in ["Items", "All"]:
                    for idx_r, row in cluster_gantt_df.iterrows():
                        prefix = "💎 [MILESTONE] " if row.get('Milestone') == True else f"[{row.get('Key', 'N/A')}]"
                        
                        epic_name_trunc = truncate_text(row['Epic Name'])
                        # Two-line format for epic
                        if view_mode == "Items":
                            epic_y_label = f"<b>{prefix}</b><br><span style='font-size:10px; color:#CBD5E1;'>{epic_name_trunc}</span>"
                        else:
                            epic_y_label = f"&nbsp;&nbsp;&nbsp;&nbsp;↳ <b>{prefix}</b><br>&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;<span style='font-size:10px; color:#CBD5E1;'>{epic_name_trunc}</span>"
                        
                        epic_id = f"E_{epic_counter}"
                            
                        r_dict = row.to_dict()
                        r_dict['Y_Axis_Label'] = epic_id
                        r_dict['Display_Label'] = epic_y_label
                        r_dict['Is_Cluster_Header'] = False
                        r_dict['Y_Index'] = epic_counter
                        plotting_rows.append(r_dict)
                        y_axis_order.append(epic_id)
                        epic_counter += 1
                        
            if plotting_rows:
                plot_df = pd.DataFrame(plotting_rows)
                
                # Render Plotly timeline using numeric Y axis to prevent auto-grouping
                fig = px.timeline(
                    plot_df,
                    x_start="Start Date",
                    x_end="Due Date",
                    y="Y_Index", # NUMERIC AXIS ensures absolute placement
                    color="Gantt_Status",
                    color_discrete_map=status_colors,
                    category_orders={
                        "Gantt_Status": ["To Do", "In Progress", "Done", "Blocked", "Milestone", "Cluster Header"]
                    },
                    custom_data=["Key", "Epic Name", "Status", "Quarter", "Sprint_Visual", "Cluster", "Start Date", "Due Date", "Is_Cluster_Header", "Size"]
                )
                
                # Configure Y-axis: hide native ticks and use left-aligned annotations
                # Altura ampliada por tener dos filas por item
                dynamic_height = max(280, 100 + len(plot_df) * 55)
                fig.update_yaxes(
                    autorange="reversed", 
                    title_text="", 
                    showticklabels=False,
                    tickmode="array",
                    tickvals=plot_df['Y_Index'].tolist()
                )
                fig.update_xaxes(title_text="", tickfont=dict(color="#94A3B8"), range=shared_xaxis_range)
                
                # Add vertical indicator line for today's date
                today_dt = pd.Timestamp.today()
                fig.add_vline(x=today_dt, line_width=2, line_dash="dash", line_color="#EF4444")
                fig.add_annotation(x=today_dt, y=1.0, yref="paper", text="<b>📍 Today</b>", showarrow=False, font=dict(color="#EF4444", size=11), xanchor="center", yanchor="bottom")
                
                # Draw vertical lines indicating Quarter boundaries
               
                if not quarters_sel:
                    for q_month, q_name in [(4, "Q2"), (7, "Q3"), (10, "Q4")]:
                        q_date = pd.Timestamp(year=año_activo, month=q_month, day=1)
                        fig.add_vline(x=q_date, line_width=1.5, line_dash="dot", line_color="#475569")
                        fig.add_annotation(x=q_date, y=1.01, yref="paper", text=f"<b>{q_name}</b>", showarrow=False, font=dict(color="#94A3B8", size=10), xanchor="left")
    
                # Set custom tooltip values and constant bar widths
                fig.update_traces(
                    width=0.7,  # Unify bar widths to prevent overlapping
                    hovertemplate=(
                        "<b>🏷️ Epic/US:</b> [%{customdata[0]}] <b>%{customdata[1]}</b><br>"
                        "<b>📏 Size:</b> %{customdata[9]}<br>"
                        "<b>📌 Status:</b> %{customdata[2]}<br>"
                        "<b>📅 Planning:</b> %{customdata[3]} | %{customdata[4]}<br>"
                        "<b>💼 Cluster:</b> %{customdata[5]}<br>"
                        "<b>🕒 Dates:</b> %{customdata[6]|%d-%m-%Y} to %{customdata[7]|%d-%m-%Y}<extra></extra>"
                    ),
                    marker_line_width=1.5, marker_line_color="white", opacity=0.95
                )
                
                # Customize Plotly chart layout, gridlines, and backgrounds
                fig.update_layout(
                    barmode="overlay", # Center bars precisely on their Y-indices
                    barcornerradius=8, clickmode="event+select",
                    plot_bgcolor="rgba(0,0,0,0)", paper_bgcolor="rgba(0,0,0,0)",
                    height=dynamic_height,
                    margin=dict(l=450, r=10, t=35, b=10), # Configure chart borders and dimensions
                    xaxis=dict(showgrid=True, gridcolor="#334155"),
                    yaxis=dict(showgrid=False),
                    showlegend=True, 
                    legend_title_text="",
                    # Position interactive legend inline at top-left
                    legend=dict(orientation="h", yanchor="bottom", y=1.01, xanchor="left", x=-0.4, font=dict(color="#E2E8F0"))
                )
                
                # Draw left-aligned visual text labels mimicking a desktop tree list
                for idx, row in plot_df.iterrows():
                    fig.add_annotation(
                        xref="paper",
                        x=0,
                        xshift=-440,  # Align label leftwards within the margin
                        y=row['Y_Index'], # Anchor the text annotation vertically to the row index
                        yref="y",
                        text=row['Display_Label'],
                        showarrow=False,
                        xanchor="left",
                        align="left",
                        font=dict(color="#F8FAFC", size=11, family="Inter, sans-serif")
                    )
                
                # Render chart with unique key to guarantee reactivity
                chart_key = f"gantt_chart_main_{st.session_state.chart_key_counter}"
                select_event = st.plotly_chart(fig, use_container_width=True, on_select="rerun", key=chart_key)
                
                # Handle selection events on bar clicks
                if select_event and "selection" in select_event and select_event["selection"]:
                    points = select_event["selection"].get("points", [])
                    if len(points) > 0:
                        c_data = points[0].get("customdata")
                        if c_data:
                            clicked_key = c_data[0]
                            cluster_nm = c_data[5]
                            is_header = c_data[8]
                            
                            if is_header:
                                # Ignore selection events on Cluster headers
                                st.session_state.chart_key_counter += 1
                                st.rerun()
                            else:
                                # Sync selection if an individual Epic bar is clicked
                                if clicked_key != st.session_state.selected_epic_key:
                                    st.session_state.selected_epic_key = clicked_key
                                    epic_name = c_data[1]
                                    st.session_state["selectbox_selected_epic"] = f"[{clicked_key}] {epic_name}"
                                    st.session_state.chart_key_counter += 1
                                    st.rerun()
            else:
                st.info("⚠️ No epics match the selected filters.")
        else:
            st.info("⚠️ No epics match the selected filters.")

    # 8. Expanded Selected Epic Details card
    st.divider()
    detail_expanded = st.session_state.selected_epic_key is not None
    with st.expander("🔍 Check Items details", expanded=detail_expanded):
        st.markdown('<div class="epic-details-accent"></div>', unsafe_allow_html=True)
        if not filtered_df.empty:
            filtered_df['Select_Option'] = filtered_df.apply(make_select_option, axis=1)
            options_list = sorted(list(filtered_df["Select_Option"].unique()))
            
            # Initialize default key on first load
            if not st.session_state.selected_epic_key and len(options_list) > 0:
                first_opt = options_list[0]
                match_key = re.match(r'^\[(.*?)\]', first_opt)
                if match_key:
                    st.session_state.selected_epic_key = match_key.group(1)
                    st.session_state["selectbox_selected_epic"] = first_opt
            
            # Sync key to selectbox session state
            elif st.session_state.selected_epic_key:
                match_opt = [opt for opt in options_list if opt.startswith(f"[{st.session_state.selected_epic_key}]")]
                if match_opt:
                    st.session_state["selectbox_selected_epic"] = match_opt[0]
            
            selected_option = st.selectbox(
                "Select an Epic/US", 
                options_list, 
                key="selectbox_selected_epic",
                help="You can click directly on any bar in the Gantt chart above to view its details instantly."
            )
            
            # Manual sync on selectbox selection
            if selected_option:
                match_key = re.match(r'^\[(.*?)\]', selected_option)
                if match_key:
                    manual_key = match_key.group(1)
                    if manual_key != st.session_state.selected_epic_key:
                        st.session_state.selected_epic_key = manual_key
                        st.rerun()
            
            epic_data = filtered_df[filtered_df["Select_Option"] == selected_option].iloc[0]
            
            # Dynamic Status Accent Color Mapping for Selected Epic Details
            status_val_raw = clean_val(epic_data.get('Status'), 'To Do')
            
            status_group = get_status_group(status_val_raw, epic_data.get('Milestone') == True)
            
            status_colors_accent = {
                "To Do": "#94A3B8",       # Light Gray
                "In Progress": "#60A5FA",  # Light Blue
                "Done": "#4ADE80",        # Light Green
                "Blocked": "#F87171",     # Light Red
                "Milestone": "#FACC15"    # Light Yellow
            }
            status_color = status_colors_accent.get(status_group, "#60A5FA")
            
            # Formulate clear status display with both group and raw status
            status_display = status_group
            if status_val_raw.lower() != status_group.lower() and status_val_raw != "N/A":
                status_display = f"{status_group} ({status_val_raw})"
            
            # Inject dynamic background and border color for the Selected Epic Details expander box below the Gantt chart
            st.markdown(f"""
                <style>
                    div[data-testid="stExpander"]:has(.epic-details-accent) {{
                        background-color: #18181D !important;
                        border: 1px solid #3E3E4A !important;
                        box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.4) !important;
                        transition: all 0.3s ease-in-out !important;
                    }}
                    div[data-testid="stExpander"]:has(.epic-details-accent) > details,
                    div[data-testid="stExpander"]:has(.epic-details-accent) summary {{
                        background-color: #18181D !important;
                    }}
                    div[data-testid="stExpander"]:has(.epic-details-accent) summary:hover {{
                        background-color: #2D2D38 !important;
                    }}
                </style>
            """, unsafe_allow_html=True)

            col1, col2, col3 = st.columns(3)
            with col1:
                st.markdown(f"**Jira Key:** `{clean_val(epic_data.get('Key'), 'N/A')}`")
                st.markdown(f"**Status:** <span style='background-color: {status_color}22; color: {status_color}; border: 1px solid {status_color}; padding: 3px 10px; border-radius: 12px; font-size: 13px; font-weight: 600; display: inline-block;'>{status_display}</span>", unsafe_allow_html=True)
            with col2:
                sprint_raw = epic_data.get('Sprint')
                sprint_display = "N/A"
                if pd.notna(sprint_raw) and str(sprint_raw).strip().lower() != "nan" and sprint_raw != "":
                    try:
                        sprint_display = str(int(float(sprint_raw)))
                    except:
                        sprint_display = str(sprint_raw)
                st.markdown(f"**Sprint:** {sprint_display}")
                epic_sz = clean_val(epic_data.get('Size'), 'M')
                st.markdown(f"**Size:** <span style='background-color: #60A5FA22; color: #60A5FA; border: 1px solid #60A5FA; padding: 2px 8px; border-radius: 8px; font-size: 12px; font-weight: 600; display: inline-block;'>{epic_sz}</span>", unsafe_allow_html=True)
                st.markdown(f"**Cluster:** {clean_val(epic_data.get('Cluster'), 'N/A')}")
            with col3:
                start_str = epic_data['Start Date'].strftime('%d %b %Y') if pd.notna(epic_data.get('Start Date')) else "Unassigned"
                due_str = epic_data['Due Date'].strftime('%d %b %Y') if pd.notna(epic_data.get('Due Date')) else "Unassigned"
                st.markdown(f"**Start Date:** {start_str}")
                st.markdown(f"**End Date:** {due_str}")
                
            st.markdown("---")
            st.markdown("**Description:**")
            st.info(clean_val(epic_data.get('Description'), 'No description provided.'))
        else:
            st.info("No selectable epics found with active filters.")

    # Download buttons below the Gantt & Details
    st.divider()
    st.subheader("💾 Export Roadmap")
    col_pdf, col_csv = st.columns(2)
    
    with col_pdf:
        # Build PDF bytes
        pdf_data = build_quarterly_plan_pdf(
            st.session_state.main_df,
            st.session_state.get("primary_color", "#0B2756")
        )
        st.download_button(
            label="⬇️ Download Plan (PDF)",
            data=pdf_data,
            file_name="quarterly_plan_roadmap.pdf",
            mime="application/pdf",
            use_container_width=True
        )
        
    with col_csv:
        export_df = st.session_state.main_df.copy()
        if 'Start Date' in export_df.columns: 
            export_df['Start Date'] = pd.to_datetime(export_df['Start Date']).dt.strftime('%d-%m-%Y')
        if 'Due Date' in export_df.columns: 
            export_df['Due Date'] = pd.to_datetime(export_df['Due Date']).dt.strftime('%d-%m-%Y')
            
        csv_data = export_df.to_csv(index=False).encode('utf-8')
        st.download_button(
            label="⬇️ Download Plan (CSV)", 
            data=csv_data, 
            file_name="epics_sprints_jira_edited.csv", 
            mime="text/csv",
            use_container_width=True
        )

# ---------------------------------------------------------
# STEP NAVIGATION EXECUTION
# ---------------------------------------------------------
if st.session_state.active_tab_qp == "💾 Exporter":
    progress_df = st.session_state.get("quarterly_progress_df")
    if isinstance(progress_df, pd.DataFrame) and not progress_df.empty:
        st.subheader("📈 Quarterly Epic Progress Slide")
        st.write("This presentation-ready slide uses the same Epic progress fields as Sprint Review, with your additional update/context column.")
        progress_title = st.session_state.get("quarterly_progress_config", {}).get("title", "Quarterly Epic Progress")
        progress_subtitle = st.session_state.get("quarterly_progress_config", {}).get("subtitle", "")
        roadmap_title = st.session_state.get("quarterly_progress_config", {}).get("roadmap_title", "Delivery Roadmap")
        roadmap_subtitle = st.session_state.get("quarterly_progress_config", {}).get("roadmap_subtitle", "")
        progress_pdf = build_quarterly_progress_slide_pdf(
            prepare_quarterly_progress_for_presentation(progress_df),
            progress_title,
            st.session_state.get("primary_color", "#0B2756"),
        )
        st.download_button(
            label="⬇️ Download Quarterly Epic Progress Slide (PDF)",
            data=progress_pdf,
            file_name="quarterly_epic_progress.pdf",
            mime="application/pdf",
            use_container_width=True,
        )
        st.markdown("#### 👁️ Progress slide preview")
        progress_preview = base64.b64encode(progress_pdf).decode("utf-8")
        st.markdown(
            f'<iframe src="data:application/pdf;base64,{progress_preview}" width="100%" height="520" '
            'type="application/pdf" style="border: 1px solid #3E3E4A; border-radius: 12px; background: #ffffff;"></iframe>',
            unsafe_allow_html=True,
        )
        st.markdown("### 🗓️ Delivery roadmap timeline")
        roadmap_df, roadmap_timeline = build_delivery_roadmap_timeline(progress_df)
        if roadmap_timeline is None:
            st.info("Select at least one Epic in Workbook and set its delivery month to build the client delivery timeline.")
        else:
            st.plotly_chart(roadmap_timeline, use_container_width=True)
            st.caption(f"Showing {len(roadmap_df)} selected delivery milestone(s). Hover over a milestone to see its completion, status, and context.")
            roadmap_pdf = build_delivery_roadmap_slide_pdf(
                roadmap_df,
                f"{progress_title} - Delivery Roadmap",
                st.session_state.get("primary_color", "#0B2756"),
            )
            st.download_button(
                label="⬇️ Download Delivery Roadmap Slide (PDF)",
                data=roadmap_pdf,
                file_name="quarterly_delivery_roadmap.pdf",
                mime="application/pdf",
                use_container_width=True,
            )
            powerpoint_data = build_quarterly_progress_pptx(
                prepare_quarterly_progress_for_presentation(progress_df),
                roadmap_df,
                progress_title,
                progress_subtitle,
                roadmap_title,
                roadmap_subtitle,
                st.session_state.get("primary_color", "#0B2756"),
            )
            st.download_button(
                label="⬇️ Download PowerPoint (Progress + Roadmap)",
                data=powerpoint_data,
                file_name="quarterly_epic_progress_and_roadmap.pptx",
                mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                use_container_width=True,
            )
            st.markdown("#### 👁️ Delivery roadmap slide preview")
            roadmap_preview = base64.b64encode(roadmap_pdf).decode("utf-8")
            st.markdown(
                f'<iframe src="data:application/pdf;base64,{roadmap_preview}" width="100%" height="520" '
                'type="application/pdf" style="border: 1px solid #3E3E4A; border-radius: 12px; background: #ffffff;"></iframe>',
                unsafe_allow_html=True,
            )
        st.divider()

    if st.session_state.get("main_df") is not None:
        show_timeline_gantt_tab()
    elif not (isinstance(progress_df, pd.DataFrame) and not progress_df.empty):
        st.warning("⚠️ Load a backlog or Quarterly Epic Progress in Ingestion first.")
