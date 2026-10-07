import pandas as pd
import re
import os
import io
import html
import plotly.express as px
# Helper to parse Quarter (e.g. "Q1 - 26") to a starting Date
def parse_quarter_to_date(q_str):
    if pd.isna(q_str) or not isinstance(q_str, str): return pd.Timestamp.today()
    match = re.search(r'Q([1-4]).*?(\d{2,4})', q_str, re.IGNORECASE)
    if match:
        q = int(match.group(1))
        year = int(match.group(2))
        if year < 100: year += 2000
        month = (q - 1) * 3 + 1
        return pd.Timestamp(year=year, month=month, day=1)
    return pd.Timestamp.today()


def get_status_group(status_str, is_milestone=False):
    """
    Unifies status mapping across the Gantt chart and selected detail panel views.
    """
    if is_milestone:
        return "Milestone"
    s = str(status_str).strip().lower()
    if s in ["to do", "open", "backlog", "pending", "por hacer", "abierto"]:
        return "To Do"
    elif s in ["in progress", "en progreso", "desarrollo", "testing", "in_progress"]:
        return "In Progress"
    elif s in ["done", "completado", "listo", "closed", "cerrado", "finalizado", "terminado", "acceptance test"]:
        return "Done"
    elif s in ["blocked", "bloqueado", "impedimento", "pausado"]:
        return "Blocked"
    return "To Do"


def truncate_text(text, max_len=45):
    """
    Cleans and truncates text values for optimized rendering inside Plotly visualizations.
    """
    if pd.isna(text):
        return ""
    text_str = str(text)
    return text_str[:max_len-3] + "..." if len(text_str) > max_len else text_str


def clean_val(val, fallback="N/A"):
    """
    Sanitizes values to avoid rendering empty or 'NaN' elements in active UI layouts.
    """
    if pd.isna(val) or val == "" or str(val).strip().lower() == "nan":
        return fallback
    return str(val)


def make_select_option(r):
    """
    Formats a dictionary row representing a Jira epic into a clean dropdown selector option.
    """
    k = clean_val(r.get('Key'), 'N/A')
    n = clean_val(r.get('Epic Name'), 'Unnamed Epic')
    return f"[{k}] {n}"


def val_changed(v1, v2):
    """
    Safe comparison check preventing TypeError: 'boolean value of NA is ambiguous' errors in pandas.
    """
    n1, n2 = pd.isna(v1), pd.isna(v2)
    if n1 and n2:
        return False
    if n1 != n2:
        return True
    try:
        return bool(v1 != v2)
    except Exception:
        return True


def get_quarterly_team_label(labels):
    """Return the first configured team label assigned to an Epic."""
    configured_labels = [label.strip() for label in os.getenv("TEAM_LABELS", "").split(",") if label.strip()]
    item_labels = {label.strip().lower() for label in re.split(r'[\s,]+', str(labels)) if label.strip()}
    return next((label for label in configured_labels if label.lower() in item_labels), "-")


def prepare_quarterly_progress_for_presentation(df):
    """Add presentation fields and prioritize open, lower-progress Epics for review."""
    prepared = df.copy()
    presentation_defaults = {
        "Presentation update": "",
        "Include in delivery roadmap": False,
        "Release version": "",
        "Delivery month": "",
        "Delivery timing": "Mid",
    }
    for column, default_value in presentation_defaults.items():
        if column not in prepared.columns:
            prepared[column] = default_value
    prepared["Team"] = prepared.get("Labels", pd.Series("", index=prepared.index)).apply(get_quarterly_team_label)

    configured_labels = [label.strip() for label in os.getenv("TEAM_LABELS", "").split(",") if label.strip()]
    priorities = {label.lower(): position for position, label in enumerate(configured_labels)}
    prepared["_team_order"] = prepared["Team"].str.lower().map(priorities).fillna(len(priorities))
    status_values = prepared.get("Status", pd.Series("", index=prepared.index)).fillna("").astype(str).str.strip().str.lower()
    # Quarterly progress maps a closed Epic to Done.  Keep accepted variants for
    # existing imported data, then place these Epics after all open work.
    prepared["_is_closed_epic"] = status_values.isin({"done", "closed", "resolved"}).astype(int)
    prepared["_completion_order"] = pd.to_numeric(
        prepared.get("Completion", pd.Series("0", index=prepared.index))
        .fillna("0")
        .astype(str)
        .str.replace("%", "", regex=False),
        errors="coerce",
    ).fillna(0)
    prepared = prepared.sort_values(
        ["_is_closed_epic", "_team_order", "_completion_order", "Key"],
        kind="stable",
    ).drop(columns=["_is_closed_epic", "_team_order", "_completion_order"])
    return prepared.reset_index(drop=True)


def delivery_month_options():
    """Provide a simple rolling set of month choices for approximate delivery milestones."""
    start = pd.Timestamp.today().replace(day=1)
    return ["-"] + [(start + pd.DateOffset(months=offset)).strftime("%b %Y") for offset in range(18)]


def build_delivery_roadmap_timeline(df):
    """Turn selected delivery commitments into an easy-to-read milestone timeline."""
    roadmap = df[df["Include in delivery roadmap"].fillna(False).astype(bool)].copy()
    roadmap = roadmap[
        roadmap["Delivery month"].fillna("").astype(str).str.strip().ne("")
        & roadmap["Delivery month"].fillna("").astype(str).str.strip().ne("-")
    ]
    if roadmap.empty:
        return roadmap, None

    timing_days = {"Beginning": 3, "Mid": 15, "End": 25}
    def milestone_date(row):
        try:
            month_start = pd.to_datetime(f"01 {row['Delivery month']}", format="%d %b %Y")
            return month_start + pd.Timedelta(days=timing_days.get(str(row.get("Delivery timing", "Mid")), 15) - 1)
        except (TypeError, ValueError):
            return pd.NaT

    roadmap["Milestone date"] = roadmap.apply(milestone_date, axis=1)
    roadmap = roadmap.dropna(subset=["Milestone date"]).sort_values("Milestone date", kind="stable")
    if roadmap.empty:
        return roadmap, None

    roadmap["Epic"] = roadmap["Summary"].fillna("Unnamed Epic").astype(str)
    roadmap["Release"] = roadmap["Release version"].fillna("").replace("", "Release TBD")
    roadmap["Update"] = roadmap["Presentation update"].fillna("").replace("", "No additional context provided")
    fig = px.scatter(
        roadmap,
        x="Milestone date",
        y="Epic",
        symbol="Delivery timing",
        text="Release",
        custom_data=["Update", "Completion", "Status"],
    )
    fig.add_vline(x=pd.Timestamp.today(), line_width=2, line_dash="dash", line_color="#EF4444")
    fig.add_annotation(
        x=pd.Timestamp.today(), y=1.03, yref="paper", text="<b>Today</b>", showarrow=False,
        font=dict(color="#EF4444", size=11), xanchor="center"
    )
    fig.update_traces(
        marker=dict(size=17, line=dict(width=2, color="#FFFFFF")),
        textposition="top center",
        textfont=dict(size=11, color="#F8FAFC"),
        hovertemplate=(
            "<b>%{y}</b><br>"
            "<b>Target:</b> %{x|%d %b %Y}<br>"
            "<b>Release:</b> %{text}<br>"
            "<b>Completion:</b> %{customdata[1]}<br>"
            "<b>Status:</b> %{customdata[2]}<br>"
            "<b>Context:</b> %{customdata[0]}<extra></extra>"
        ),
    )
    fig.update_layout(
        title="Planned delivery milestones",
        xaxis_title="", yaxis_title="", height=max(380, 115 + len(roadmap) * 68),
        plot_bgcolor="#18181D", paper_bgcolor="#18181D", font=dict(color="#F8FAFC"),
        legend_title_text="Approx. timing", margin=dict(l=20, r=20, t=65, b=30),
    )
    fig.update_xaxes(showgrid=True, gridcolor="#475569", tickformat="%b\n%Y")
    fig.update_yaxes(showgrid=False, autorange="reversed")
    return roadmap, fig


def build_delivery_roadmap_slide_pdf(roadmap_df, title, primary_color_hex):
    """Create a compact client-facing PDF slide with delivery milestones on a time axis."""
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib import colors
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    width, height = landscape(letter)
    pdf = canvas.Canvas(buffer, pagesize=(width, height))
    primary = colors.HexColor(primary_color_hex)
    background = colors.HexColor("#F8FAFC")
    dark = colors.HexColor("#1E293B")
    muted = colors.HexColor("#64748B")

    pdf.setFillColor(background)
    pdf.rect(0, 0, width, height, stroke=0, fill=1)
    pdf.setFillColor(primary)
    pdf.rect(0, height - 18, width, 18, stroke=0, fill=1)
    pdf.rect(0, 0, 8, height, stroke=0, fill=1)
    pdf.rect(width - 8, 0, 8, height, stroke=0, fill=1)

    pdf.setFillColor(primary)
    pdf.setFont("Helvetica-Bold", 24)
    pdf.drawString(54, height - 58, str(title)[:85])
    pdf.setFillColor(muted)
    pdf.setFont("Helvetica", 10)
    pdf.drawString(54, height - 76, "Committed Epic delivery roadmap - planned milestones for unfinished work")

    if roadmap_df.empty:
        pdf.setFillColor(dark)
        pdf.setFont("Helvetica", 14)
        pdf.drawString(54, height - 130, "No delivery milestones selected.")
        pdf.save()
        buffer.seek(0)
        return buffer.getvalue()

    earliest = roadmap_df["Milestone date"].min().replace(day=1)
    latest = (roadmap_df["Milestone date"].max().replace(day=1) + pd.DateOffset(months=1))
    if latest <= earliest:
        latest = earliest + pd.DateOffset(months=1)
    total_days = max((latest - earliest).days, 1)
    axis_left, axis_right, axis_y = 280, width - 58, height - 130

    def x_for_date(value):
        return axis_left + ((value - earliest).days / total_days) * (axis_right - axis_left)

    pdf.setStrokeColor(colors.HexColor("#94A3B8"))
    pdf.setLineWidth(2)
    pdf.line(axis_left, axis_y, axis_right, axis_y)
    month = earliest
    while month <= latest:
        x_pos = x_for_date(month)
        pdf.setStrokeColor(colors.HexColor("#CBD5E1"))
        pdf.setLineWidth(0.8)
        pdf.line(x_pos, 66, x_pos, axis_y + 10)
        pdf.setFillColor(muted)
        pdf.setFont("Helvetica-Bold", 8)
        pdf.drawCentredString(x_pos, axis_y + 16, month.strftime("%b %Y"))
        month += pd.DateOffset(months=1)

    today_x = x_for_date(pd.Timestamp.today())
    if axis_left <= today_x <= axis_right:
        pdf.setStrokeColor(colors.HexColor("#EF4444"))
        pdf.setDash(4, 3)
        pdf.line(today_x, 58, today_x, axis_y + 28)
        pdf.setDash()
        pdf.setFillColor(colors.HexColor("#EF4444"))
        pdf.setFont("Helvetica-Bold", 8)
        pdf.drawCentredString(today_x, axis_y + 30, "TODAY")

    available_height = axis_y - 70
    row_height = min(46, max(25, available_height / max(len(roadmap_df), 1)))
    max_rows = int(available_height // row_height)
    display_rows = roadmap_df.head(max_rows)
    for index, (_, row) in enumerate(display_rows.iterrows()):
        y_pos = axis_y - 33 - index * row_height
        marker_x = x_for_date(row["Milestone date"])
        pdf.setStrokeColor(colors.HexColor("#CBD5E1"))
        pdf.setLineWidth(0.5)
        pdf.line(54, y_pos - 13, axis_right, y_pos - 13)
        from reportlab.lib.utils import simpleSplit
        pdf.setFillColor(dark)
        pdf.setFont("Helvetica-Bold", 7.5)
        epic_label = str(row.get("Summary", "Unnamed Epic"))
        epic_lines = simpleSplit(epic_label, "Helvetica-Bold", 7.5, 215)[:3]
        for line_index, line in enumerate(epic_lines):
            pdf.drawString(54, y_pos + 6 - line_index * 8, line)
        release = row.get("Release", "Release TBD")
        pdf.setStrokeColor(primary)
        pdf.setLineWidth(1.2)
        pdf.line(marker_x, y_pos + 6, marker_x, axis_y)
        pdf.setFillColor(primary)
        pdf.circle(marker_x, y_pos + 6, 5, stroke=0, fill=1)
        pdf.setFillColor(muted)
        pdf.setFont("Helvetica-Bold", 7.5)
        pdf.drawString(marker_x + 8, y_pos + 3, str(release))

    if len(roadmap_df) > len(display_rows):
        pdf.setFillColor(muted)
        pdf.setFont("Helvetica-Oblique", 8)
        pdf.drawRightString(axis_right, 46, f"+ {len(roadmap_df) - len(display_rows)} additional milestone(s) shown in the interactive timeline")
    pdf.setFillColor(muted)
    pdf.setFont("Helvetica", 8)
    pdf.drawString(54, 32, f"Generated {pd.Timestamp.today().strftime('%d %b %Y')} | Delivery dates are approximate (Beginning, Mid, End of month).")
    pdf.save()
    buffer.seek(0)
    return buffer.getvalue()


def build_quarterly_progress_slide_pdf(df, title, primary_color_hex):
    """Create a presentation-style landscape slide for the quarterly Epic progress."""
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle
    from utils.pdf_helpers import get_progress_bar_drawing

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(letter),
        leftMargin=54, rightMargin=54, topMargin=48, bottomMargin=42
    )
    styles = getSampleStyleSheet()
    primary = colors.HexColor(primary_color_hex)
    title_style = ParagraphStyle(
        "QuarterlyProgressTitle", parent=styles["Normal"], fontName="Helvetica-Bold",
        fontSize=25, leading=30, textColor=primary, spaceAfter=5
    )
    subtitle_style = ParagraphStyle(
        "QuarterlyProgressSubtitle", parent=styles["Normal"], fontName="Helvetica",
        fontSize=10, leading=13, textColor=colors.HexColor("#475569"), spaceAfter=16
    )
    header_style = ParagraphStyle(
        "QuarterlyProgressHeader", parent=styles["Normal"], fontName="Helvetica-Bold",
        fontSize=8.5, leading=10, textColor=colors.white
    )
    body_style = ParagraphStyle(
        "QuarterlyProgressBody", parent=styles["Normal"], fontName="Helvetica",
        fontSize=8, leading=10, textColor=colors.HexColor("#1E293B")
    )
    key_style = ParagraphStyle("QuarterlyProgressKey", parent=body_style, fontName="Helvetica-Bold")

    visible_columns = [
        ("Key", 55), ("Epic", 150), ("Progress", 60), ("Closed / Resolved", 56),
        ("To Do", 45), ("In progress", 58), ("Status", 62), ("Team", 75), ("Update", 123)
    ]
    column_sources = {"Epic": "Summary", "Progress": "Completion", "Update": "Presentation update"}
    table_data = [[Paragraph(name, header_style) for name, _ in visible_columns]]
    for _, row in df.iterrows():
        cells = []
        for name, _ in visible_columns:
            source = column_sources.get(name, name)
            value = row.get(source, "-")
            value = "-" if pd.isna(value) or str(value).strip() == "" else str(value)
            if name == "Progress":
                cells.append(get_progress_bar_drawing(value, primary, width=52, height=18) or Paragraph(value, body_style))
            else:
                cells.append(Paragraph(html.escape(value), key_style if name == "Key" else body_style))
        table_data.append(cells)

    table = Table(table_data, colWidths=[width for _, width in visible_columns], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), primary),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
    ]))
    project_name = os.getenv("PROJECT_NAME", "MYPROJECT")
    generated = pd.Timestamp.today().strftime("%d %b %Y")
    story = [
        Paragraph(html.escape(title), title_style),
        Paragraph(f"<b>Project:</b> {html.escape(project_name)} &nbsp; | &nbsp; <b>Generated:</b> {generated} &nbsp; | &nbsp; <b>Committed Epics:</b> {len(df)}", subtitle_style),
        table,
    ]
    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def build_quarterly_progress_pptx(progress_df, roadmap_df, progress_title, progress_subtitle, roadmap_title, roadmap_subtitle, primary_color_hex):
    """Build editable PowerPoint slides for quarterly progress and delivery milestones."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.dml.color import RGBColor

    def rgb(hex_value):
        value = hex_value.lstrip("#")
        return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))

    primary = rgb(primary_color_hex)
    navy = RGBColor(11, 39, 86)
    dark = RGBColor(30, 41, 59)
    muted = RGBColor(100, 116, 139)
    light = RGBColor(241, 245, 249)
    grid = RGBColor(203, 213, 225)
    green = RGBColor(34, 197, 94)
    template_background = RGBColor(0, 45, 55)
    template_teal = RGBColor(0, 151, 143)
    white = RGBColor(255, 255, 255)
    template_muted = RGBColor(113, 190, 187)

    template_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "assets", "templates", "quarterly_planning_template.pptx")
    using_template = os.path.exists(template_path)
    prs = Presentation(template_path) if using_template else Presentation()
    if using_template:
        for slide_id in list(prs.slides._sldIdLst):
            prs.part.drop_rel(slide_id.rId)
            prs.slides._sldIdLst.remove(slide_id)
    else:
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    content_layout = prs.slide_layouts[2] if using_template and len(prs.slide_layouts) > 2 else blank_layout

    def add_text(slide, text, left, top, width, height, size=12, color=dark, bold=False, align=PP_ALIGN.LEFT):
        shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        text_frame = shape.text_frame
        text_frame.clear()
        text_frame.word_wrap = True
        text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        paragraph = text_frame.paragraphs[0]
        paragraph.alignment = align
        run = paragraph.add_run()
        run.text = str(text)
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        return shape

    def add_chrome(slide, slide_title, subtitle=""):
        if using_template:
            # The corporate layout contains editable title placeholders.  They are
            # removed here because python-pptx does not preserve their type styling
            # reliably, causing duplicate, overlapping title text in the export.
            for shape in list(slide.placeholders):
                element = shape._element
                element.getparent().remove(element)
            add_text(slide, slide_title, 0.42, 0.62, 11.6, 0.46, size=24, color=white, bold=False)
            if subtitle:
                add_text(slide, subtitle, 0.43, 1.17, 11.4, 0.27, size=13, color=white)
            return
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(248, 250, 252)
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.18))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = primary
        top_bar.line.fill.background()
        for x_pos in [0, 13.25]:
            band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x_pos), 0, Inches(0.08), prs.slide_height)
            band.fill.solid()
            band.fill.fore_color.rgb = primary
            band.line.fill.background()
        add_text(slide, slide_title, 0.7, 0.35, 11.9, 0.46, size=25, color=navy, bold=True)
        if subtitle:
            add_text(slide, subtitle, 0.7, 0.82, 11.7, 0.24, size=10, color=muted)

    progress_columns = [
        ("Key", "Key", 0.85), ("Epic", "Summary", 2.70), ("Progress", "Completion", 0.85),
        ("Closed / Resolved", "Closed / Resolved", 0.90), ("To Do", "To Do", 0.65),
        ("In progress", "In progress", 0.90), ("Status", "Status", 0.90), ("Team", "Team", 1.10),
        ("Update", "Presentation update", 3.13),
    ]
    rows_per_slide = 12
    progress_df = prepare_quarterly_progress_for_presentation(progress_df)
    for start in range(0, len(progress_df), rows_per_slide):
        chunk = progress_df.iloc[start:start + rows_per_slide].reset_index(drop=True)
        suffix = "" if start == 0 else f" (cont. {start // rows_per_slide + 1})"
        slide = prs.slides.add_slide(content_layout)
        add_chrome(slide, f"{progress_title}{suffix}", progress_subtitle)
        table_left, table_top, table_width = 0.42, 1.90, 12.45
        table_height = 4.92 if using_template else 5.85
        row_height = table_height / (len(chunk) + 1)
        table_shape = slide.shapes.add_table(len(chunk) + 1, len(progress_columns), Inches(table_left), Inches(table_top), Inches(table_width), Inches(table_height))
        table = table_shape.table
        for row in table.rows:
            row.height = Inches(row_height)
        for index, (_, _, column_width) in enumerate(progress_columns):
            table.columns[index].width = Inches(column_width)
            cell = table.cell(0, index)
            cell.fill.solid()
            cell.fill.fore_color.rgb = template_teal if using_template else navy
            cell.text = progress_columns[index][0]
            paragraph = cell.text_frame.paragraphs[0]
            paragraph.runs[0].font.name = "Arial"
            paragraph.runs[0].font.size = Pt(8)
            paragraph.runs[0].font.bold = True
            paragraph.runs[0].font.color.rgb = white
        for row_index, (_, row) in enumerate(chunk.iterrows(), start=1):
            for col_index, (_, source, _) in enumerate(progress_columns):
                cell = table.cell(row_index, col_index)
                cell.fill.solid()
                cell.fill.fore_color.rgb = template_background if using_template else (RGBColor(255, 255, 255) if row_index % 2 else light)
                cell.margin_left = Inches(0.05)
                cell.margin_right = Inches(0.05)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                value = row.get(source, "-")
                value = "-" if pd.isna(value) or str(value).strip() == "" else str(value)
                cell.text = "" if source == "Completion" else value
                if source != "Completion":
                    paragraph = cell.text_frame.paragraphs[0]
                    paragraph.runs[0].font.name = "Arial"
                    paragraph.runs[0].font.size = Pt(7.2)
                    paragraph.runs[0].font.color.rgb = white if using_template else dark
                    if source == "Key":
                        paragraph.runs[0].font.bold = True

                if source == "Completion":
                    try:
                        percentage = max(0, min(100, int(float(value.replace("%", "")))))
                    except Exception:
                        percentage = 0
                    # Keep the percentage and its visual bar *inside* the table
                    # cell.  Separate shapes drift when PowerPoint recalculates
                    # table-row heights, while text runs remain attached to the row.
                    text_frame = cell.text_frame
                    text_frame.clear()
                    text_frame.word_wrap = False
                    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                    label = text_frame.paragraphs[0]
                    label.alignment = PP_ALIGN.CENTER
                    label_run = label.add_run()
                    label_run.text = f"{percentage}%"
                    label_run.font.name = "Arial"
                    label_run.font.size = Pt(6.5)
                    label_run.font.bold = True
                    label_run.font.color.rgb = white if using_template else dark
                    bar = text_frame.add_paragraph()
                    bar.alignment = PP_ALIGN.CENTER
                    # Five compact segments fit in the narrow Progress column in
                    # PowerPoint without wrapping onto a second line.
                    filled_blocks = int(round(percentage / 20))
                    filled_run = bar.add_run()
                    filled_run.text = "▬" * filled_blocks
                    filled_run.font.name = "Arial"
                    filled_run.font.size = Pt(7)
                    filled_run.font.color.rgb = RGBColor(170, 255, 0) if using_template and percentage == 100 else (template_teal if using_template else (green if percentage == 100 else navy))
                    remaining_run = bar.add_run()
                    remaining_run.text = "▬" * (5 - filled_blocks)
                    remaining_run.font.name = "Arial"
                    remaining_run.font.size = Pt(7)
                    remaining_run.font.color.rgb = RGBColor(0, 98, 110) if using_template else RGBColor(226, 232, 240)

    if roadmap_df is not None and not roadmap_df.empty:
        slide = prs.slides.add_slide(content_layout)
        add_chrome(slide, roadmap_title, roadmap_subtitle)
        earliest = roadmap_df["Milestone date"].min().replace(day=1)
        latest = roadmap_df["Milestone date"].max().replace(day=1) + pd.DateOffset(months=1)
        total_days = max((latest - earliest).days, 1)
        axis_left, axis_right, axis_y = 0.75, 12.55, 4.55

        def x_for_date(value):
            return axis_left + ((value - earliest).days / total_days) * (axis_right - axis_left)

        axis = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(axis_left), Inches(axis_y), Inches(axis_right - axis_left), Inches(0.025))
        axis.fill.solid()
        axis.fill.fore_color.rgb = white if using_template else muted
        axis.line.fill.background()
        month = earliest
        while month <= latest:
            x_pos = x_for_date(month)
            grid_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x_pos), Inches(3.95), Inches(0.012), Inches(0.68))
            grid_line.fill.solid()
            grid_line.fill.fore_color.rgb = template_teal if using_template else grid
            grid_line.line.fill.background()
            add_text(slide, month.strftime("%b %Y"), x_pos - 0.52, 3.58, 1.04, 0.20, size=11 if using_template else 8, color=template_muted if using_template else muted, bold=True, align=PP_ALIGN.CENTER)
            month += pd.DateOffset(months=1)

        today_x = x_for_date(pd.Timestamp.today())
        if axis_left <= today_x <= axis_right:
            today_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(today_x), Inches(2.0), Inches(0.018), Inches(2.65))
            today_line.fill.solid()
            today_line.fill.fore_color.rgb = RGBColor(239, 68, 68)
            today_line.line.fill.background()
            add_text(slide, "TODAY", today_x - 0.27, 1.83, 0.56, 0.16, size=7, color=RGBColor(239, 68, 68), bold=True, align=PP_ALIGN.CENTER)

        # One release card per milestone mirrors the supplied Release Calendar:
        # release version first, followed by the Epics planned for that release.
        displayed = roadmap_df.head(12).copy()
        displayed["Milestone date"] = pd.to_datetime(displayed["Milestone date"])
        release_cards = []
        for (milestone, release), rows in displayed.groupby(["Milestone date", "Release"], sort=True):
            release_cards.append({
                "date": milestone,
                "release": str(release or "Release TBD"),
                "epics": [
                    f"{str(key).strip()} - {str(summary).strip()}" if str(key).strip() else str(summary).strip()
                    for key, summary in zip(rows.get("Key", pd.Series("", index=rows.index)), rows["Summary"])
                    if str(summary).strip()
                ],
            })

        for index, card in enumerate(release_cards):
            marker_x = x_for_date(card["date"])
            box_width = 2.45
            box_height = min(1.25, 0.42 + (0.23 * min(len(card["epics"]), 3)))
            box_left = max(0.55, min(12.70 - box_width, marker_x - (box_width / 2)))
            box_top = 2.05 if index % 2 else 2.88
            card_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(box_left), Inches(box_top), Inches(box_width), Inches(box_height))
            card_shape.fill.solid()
            card_shape.fill.fore_color.rgb = template_background if using_template else RGBColor(255, 255, 255)
            card_shape.line.color.rgb = template_teal if using_template else primary
            add_text(slide, card["release"], box_left + 0.13, box_top + 0.08, box_width - 0.25, 0.18, size=9, color=template_teal if using_template else primary, bold=True)
            epic_lines = "\n".join(f"•  {epic}" for epic in card["epics"][:3])
            if len(card["epics"]) > 3:
                epic_lines += f"\n•  +{len(card['epics']) - 3} more"
            add_text(slide, epic_lines, box_left + 0.13, box_top + 0.27, box_width - 0.25, box_height - 0.32, size=7.7, color=white if using_template else dark)
            connector = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(marker_x), Inches(box_top + box_height), Inches(0.018), Inches(max(0.02, axis_y - (box_top + box_height))))
            connector.fill.solid()
            connector.fill.fore_color.rgb = template_teal if using_template else primary
            connector.line.fill.background()
            marker = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(marker_x - 0.09), Inches(axis_y - 0.09), Inches(0.18), Inches(0.18))
            marker.fill.solid()
            marker.fill.fore_color.rgb = template_teal if using_template else primary
            marker.line.fill.background()
        marker = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(0.70), Inches(6.42), Inches(0.14), Inches(0.14))
        marker.fill.solid()
        marker.fill.fore_color.rgb = template_teal if using_template else primary
        marker.line.fill.background()
        add_text(slide, "Release milestone", 0.92, 6.39, 1.5, 0.20, size=7.5, color=white if using_template else dark, bold=True)
        add_text(slide, f"Generated {pd.Timestamp.today().strftime('%d %b %Y')} - delivery dates are approximate", 0.48, 6.95, 5.5, 0.16, size=7.5, color=template_muted if using_template else muted)

    output = io.BytesIO()
    prs.save(output)
    output.seek(0)
    return output.getvalue()


def build_quarterly_plan_pdf(df, primary_color_hex):
    import io
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors

    pdf_buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        pdf_buffer,
        pagesize=landscape(letter),
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=22,
        textColor=colors.HexColor(primary_color_hex),
        spaceAfter=12
    )
    
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        textColor=colors.HexColor("#64748B"),
        spaceAfter=20
    )
    
    cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=11
    )
    
    cell_bold_style = ParagraphStyle(
        'TableCellBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=11
    )

    story = []
    story.append(Paragraph("🎯 QUARTERLY PLAN ROADMAP", title_style))
    proj_name = os.getenv("PROJECT_NAME", "MYPROJECT")
    today_str = pd.Timestamp.today().strftime('%d-%b-%Y')
    story.append(Paragraph(f"<b>Project:</b> {proj_name} | <b>Generated:</b> {today_str} | <b>Total Epics:</b> {len(df)}", subtitle_style))
    
    headers = ["Key", "Epic Name", "Status", "Sprint", "Size", "Quarter", "Cluster"]
    table_data = [[Paragraph(f"<b>{h}</b>", cell_bold_style) for h in headers]]
    
    sorted_df = df.copy()
    if 'Quarter' in sorted_df.columns:
        sorted_df = sorted_df.sort_values(by=['Quarter', 'Cluster', 'Sprint'], na_position='last')
        
    for _, row in sorted_df.iterrows():
        sprint_val = row.get("Sprint")
        sprint_str = f"Sprint {int(float(sprint_val))}" if pd.notna(sprint_val) and str(sprint_val).strip() != "" and str(sprint_val).strip().lower() != "nan" else "-"
        
        row_cells = [
            Paragraph(str(row.get("Key", "-")), cell_bold_style),
            Paragraph(str(row.get("Epic Name", "-")), cell_style),
            Paragraph(str(row.get("Status", "-")), cell_style),
            Paragraph(sprint_str, cell_style),
            Paragraph(str(row.get("Size", "-")), cell_style),
            Paragraph(str(row.get("Quarter", "-")), cell_style),
            Paragraph(str(row.get("Cluster", "-")), cell_style),
        ]
        table_data.append(row_cells)
        
    col_widths = [60, 260, 80, 80, 50, 90, 100]
    t = Table(table_data, colWidths=col_widths, repeatRows=1)
    
    t_style = TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor(primary_color_hex)),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ('TOPPADDING', (0,0), (-1,-1), 6),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.HexColor("#FFFFFF"), colors.HexColor("#F8FAFC")]),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
    ])
    
    for i in range(len(headers)):
        table_data[0][i].style.textColor = colors.white
        
    t.setStyle(t_style)
    story.append(t)
    
    doc.build(story)
    pdf_buffer.seek(0)
    return pdf_buffer


